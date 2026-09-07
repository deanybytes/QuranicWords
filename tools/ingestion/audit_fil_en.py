import openpyxl, re

wb = openpyxl.load_workbook('db/qw_en.xlsx', data_only=True)
sheet = wb['Fil']

stop_words = {'and', 'so', 'then', 'but', 'or', 'they', 'we', 'who', 'that', 'it', 'is', 'are', 'he', 'she', 'you', 'i', 'the', 'a', 'an', 'to', 'of', 'in', 'on', 'at', 'by', 'for', 'with', 'from'}

rows_to_fix = []
for r in range(2, sheet.max_row + 1):
    rank = sheet.cell(row=r, column=1).value
    lemma = str(sheet.cell(row=r, column=2).value or '')
    sense = str(sheet.cell(row=r, column=8).value or '')
    meaning = str(sheet.cell(row=r, column=9).value or '')
    cit = str(sheet.cell(row=r, column=10).value or '')
    t_ar = str(sheet.cell(row=r, column=11).value or '')
    t_en = str(sheet.cell(row=r, column=13).value or '').strip()
    v_en = str(sheet.cell(row=r, column=14).value or '')
    
    if t_en.lower() in stop_words or f'([{t_en}])' not in v_en:
        clean_v = re.sub(r'\(\[([^\]]+)\]\)', r'\1', v_en)
        rows_to_fix.append((r, rank, lemma, meaning, cit, t_ar, t_en, clean_v))

print(f'Total rows needing target English review: {len(rows_to_fix)}')
for r, rank, lemma, meaning, cit, t_ar, t_en, clean_v in rows_to_fix[:15]:
    print(f'Row {r:4d} | Rank {rank:4d} | {lemma} ({meaning}) | Cit={cit} | CurTarget={t_en}')
    print(f'   Verse: {clean_v[:80]}...')
