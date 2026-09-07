import openpyxl, re, json

with open('db/corpus_data/quran_uthmani.json') as f:
    verses_list = json.load(f)['verses']
    uth_map = {v['verse_key']: v['text_uthmani'] for v in verses_list}

WAQF_CHARS = set('ۛۖۗۚۘۙۜ۩۞')

def get_words(vk):
    raw = uth_map.get(vk, '').split()
    return [w for w in raw if not (len(w) == 1 and w in WAQF_CHARS)]

def parse_morphology():
    corpus = {}
    with open('db/corpus_data/quranic-corpus-morphology-0.4.txt') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#') or line.startswith('LOCATION'):
                continue
            parts = line.split('\t')
            if len(parts) < 4: continue
            loc = parts[0].strip('()')
            s, v, w, seg = map(int, loc.split(':'))
            key = (s, v, w)
            if key not in corpus:
                corpus[key] = []
            corpus[key].append((parts[1], parts[2], parts[3]))
    return corpus

corpus = parse_morphology()

BW_MAP = {
    "'": 'ء', '>': 'أ', '<': 'إ', '|': 'آ', '{': 'ٱ', '&': 'ؤ', '}': 'ئ',
    'A': 'ا', 'b': 'ب', 'p': 'ة', 't': 'ت', 'v': 'ث', 'j': 'ج', 'H': 'ح',
    'x': 'خ', 'd': 'د', '*': 'ذ', 'r': 'ر', 'z': 'ز', 's': 'س', '$': 'ش',
    'S': 'ص', 'D': 'ض', 'T': 'ط', 'Z': 'ظ', 'E': 'ع', 'g': 'غ', '_': 'ـ',
    'f': 'ف', 'q': 'ق', 'k': 'ك', 'l': 'ل', 'm': 'م', 'n': 'ن', 'h': 'ه',
    'w': 'و', 'Y': 'ى', 'y': 'ي', 'F': 'ً', 'N': 'ٌ', 'K': 'ٍ', 'a': 'َ',
    'u': 'ُ', 'i': 'ِ', '~': 'ّ', 'o': 'ْ', '`': 'ٰ', '^': 'ٓ'
}

def bw_to_ar(bw):
    return ''.join(BW_MAP.get(c, c) for c in bw)

def strip_tashkeel(text):
    t = str(text).replace('\u0670', 'ا').replace('\u06E5', 'و').replace('\u06E6', 'ي').replace('\u0640', '')
    t = re.sub(r'[\u064B-\u065F\u06D6-\u06ED\uFEFF]', '', t)
    t = re.sub(r'[إأآٱ]', 'ا', t)
    t = t.replace('ة', 'ه').replace('ى', 'ي')
    return t.strip()

wb = openpyxl.load_workbook('db/qw_en.xlsx', data_only=True)
sheet = wb['Fil']

for r in range(2, 35):
    lemma = str(sheet.cell(row=r, column=2).value or '')
    cit = str(sheet.cell(row=r, column=10).value or '')
    m = re.search(r'(\d+)[\:\.](\d+)', cit)
    if not m: continue
    s, v = int(m.group(1)), int(m.group(2))
    vk = f'{s}:{v}'
    words = get_words(vk)
    t_ar = str(sheet.cell(row=r, column=11).value or '')
    clean_lem = strip_tashkeel(lemma)
    
    matches = []
    for w_idx, uth_w in enumerate(words, start=1):
        segs = corpus.get((s, v, w_idx), [])
        for form, tag, feat in segs:
            if 'LEM:' in feat:
                m_lem = re.search(r'LEM:([^\|]+)', feat)
                if m_lem:
                    corp_lem = bw_to_ar(m_lem.group(1))
                    if strip_tashkeel(corp_lem) == clean_lem:
                        has_prefix = False
                        p_char = ''
                        clean_w = uth_w
                        for p_form, p_tag, p_feat in segs:
                            if any(p in p_feat for p in ['PREFIX|w:', 'PREFIX|w+', 'PREFIX|f:', 'PREFIX|f+']):
                                has_prefix = True
                                p_char = 'و' if 'w' in p_feat else 'ف'
                        matches.append((w_idx, uth_w, has_prefix, p_char))
                        break
    print(f'Row {r:2d}: Lemma={lemma:<8} | Cit={vk:<6} | OldTarget={t_ar:<10} | CorpusMatches={matches}')
