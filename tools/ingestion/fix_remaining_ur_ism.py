import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ur.xlsx')
ws = wb['Ism']

FIXES_4 = {
    ('Al-Baqarah 2:68', 'ارِضٌ'): 'بوڑھا',
    ('Al-Baqarah 2:234', 'خَبِيرٌ'): 'واقف ہے',
    ('Al-Baqarah 2:17', 'ظُلُمَـٰتٍ'): 'اندھیروں',
    ('An-Nisa 4:78', 'يَفْقَهُونَ'): 'سمجھ سکتے'
}

count = 0
for r in range(2, ws.max_row + 1):
    cit = ws.cell(r, 10).value
    tar_ar = ws.cell(r, 11).value
    key = (cit, tar_ar)
    if key in FIXES_4:
        phrase = FIXES_4[key]
        ws.cell(r, 13).value = phrase
        v_ur = str(ws.cell(r, 14).value or '')
        clean_v = re.sub(r'\(\[([^\]]+)\]\)', r'\1', v_ur)
        if phrase in clean_v:
            ws.cell(r, 14).value = clean_v.replace(phrase, f'([{phrase}])', 1)
            count += 1
        else:
            print(f'Still failed in row {r} for phrase \"{phrase}\"')

print(f'Applied remaining fixes to {count} rows in Urdu Ism.')
wb.save('db/qw_ur.xlsx')
print('Saved db/qw_ur.xlsx.')
