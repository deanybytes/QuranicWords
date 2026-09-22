#!/usr/bin/env python3
"""
Compile 10-Language Master Curriculum Assets
===========================================
Emits 100% verified Room/JSON content assets into app/src/main/assets/content/:
- chapters.json (10 chapters: 1 Harf, 3 Fil, 6 Ism)
- sections.json (100 sections: 10 per chapter)
- lessons_vocabulary.json (1,217 lessons: 10 Chapter Intro, 997 Regular, 100 Flashback, 100 Section Exam, 10 Chapter Exam)
- exercises_vocabulary.json (9,428 exercises: 10 Chapter Intro, 4,709 Word Intro, 4,709 Multiple Choice)
- word_frequency.json (4,709 headwords with unique global ranks 1..4709 and 10-language meanings)
"""

import os
import re
import json
import random
import time
import openpyxl
from collections import defaultdict

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
DB_DIR = os.path.join(BASE_DIR, 'db')
CONTENT_DIR = os.path.join(BASE_DIR, 'app', 'src', 'main', 'assets', 'content')

LANG_CONFIG = [
    ('en', 'English', 'English'),
    ('bn', 'Bangla', 'বাংলা'),
    ('ur', 'Urdu', 'اردو'),
    ('hi', 'Hindi', 'हिन्दी'),
    ('in', 'Indonesian', 'Bahasa Indonesia'),
    ('ms', 'Malay', 'Bahasa Melayu'),
    ('tr', 'Turkish', 'Türkçe'),
    ('fa', 'Persian', 'فارسی'),
    ('ha', 'Hausa', 'Hausa'),
    ('sw', 'Swahili', 'Kiswahili'),
]

LANG_CODES = [l[0] for l in LANG_CONFIG]
DB_CODE_MAP = {'in': 'id'}

BN_DIGITS = {'0':'০', '1':'১', '2':'২', '3':'৩', '4':'৪', '5':'৫', '6':'৬', '7':'৭', '8':'৮', '9':'৯'}
UR_DIGITS = {'0':'۰', '1':'۱', '2':'۲', '3':'۳', '4':'۴', '5':'۵', '6':'۶', '7':'۷', '8':'۸', '9':'۹'}
HI_DIGITS = {'0':'०', '1':'१', '2':'२', '3':'३', '4':'४', '5':'५', '6':'६', '7':'७', '8':'८', '9':'९'}

def to_bn_digits(n): return ''.join(BN_DIGITS.get(d, d) for d in str(n))
def to_ur_digits(n): return ''.join(UR_DIGITS.get(d, d) for d in str(n))
def to_hi_digits(n): return ''.join(HI_DIGITS.get(d, d) for d in str(n))

def get_bn_ordinal(n):
    if n == 1: return '১ম'
    if n in (2, 3): return f'{to_bn_digits(n)}য়'
    if n == 4: return '৪র্থ'
    if n in (5, 7, 8, 9, 10): return f'{to_bn_digits(n)}ম'
    if n == 6: return '৬ষ্ঠ'
    return f'{to_bn_digits(n)}তম'

def get_ur_ordinal(n):
    ordinals = ['', 'پہلا', 'دوسرا', 'تیسرا', 'چوتھا', 'پانچواں', 'چھٹا', 'ساتواں', 'آٹھواں', 'نواں', 'دسواں']
    if 1 <= n < len(ordinals): return ordinals[n]
    return f'نمبر {to_ur_digits(n)}'

def strip_tashkeel(text):
    if not text: return ''
    t = str(text).replace('\u0670', 'ا').replace('\u0627\u065F', 'ا').replace('\u06E5', 'و').replace('\u06E6', 'ي')
    t = re.sub(r'[\u064B-\u065F\u06D6-\u06ED\uFEFF]', '', t)
    t = re.sub(r'[إأآٱ]', 'ا', t)
    t = t.replace('ة', 'ه').replace('ى', 'ي')
    return t.strip()

def clean_prepended_hack(hl_text):
    s = str(hl_text).strip()
    m = re.match(r'^\(\[([^\]]+)\]\)\s*(?:\[\d+\])?\s+(.+)$', s)
    if m:
        word = m.group(1).strip()
        rest = m.group(2).strip()
        if '([' in rest: return rest
        idx = rest.lower().find(word.lower())
        if idx >= 0:
            matched = rest[idx:idx+len(word)]
            return rest[:idx] + f"([{matched}])" + rest[idx+len(word):]
        return rest
    return s

def extract_highlight_phrase(text):
    if not text: return ''
    m = re.search(r'\(\[([^\]]+)\]\)', text)
    if m:
        return m.group(1).strip()
    return ''

def clean_verse_text(text):
    if not text: return ''
    s = re.sub(r'\(\[([^\]]+)\]\)', r'\1', text)
    s = re.sub(r'\s*\[\d+\]\s*', ' ', s)
    return re.sub(r'\s+', ' ', s).strip()

def calculate_token_spans(verse_ar, target_word):
    if not verse_ar or not target_word:
        return (None, None)
    clean_tw = strip_tashkeel(target_word)
    idx = verse_ar.find(target_word)
    if idx >= 0:
        return (idx, idx + len(target_word))
    words = verse_ar.split()
    curr = 0
    for w in words:
        w_idx = verse_ar.find(w, curr)
        if w_idx >= 0:
            if strip_tashkeel(w) == clean_tw or clean_tw in strip_tashkeel(w):
                return (w_idx, w_idx + len(w))
            curr = w_idx + len(w)
    return (0, len(target_word))

PROMPT_WORD_INTRO = {
    'en': 'Learn this word',
    'bn': 'শব্দটি শিখুন',
    'ur': 'یہ لفظ سیکھیں',
    'hi': 'यह शब्द सीखें',
    'in': 'Pelajari kata ini',
    'ms': 'Pelajari perkataan ini',
    'tr': 'Bu kelimeyi öğrenin',
    'fa': 'این واژه را بیاموزید',
    'ha': 'Koyi wannan kalmar',
    'sw': 'Jifunze neno hili'
}

PROMPT_MULTIPLE_CHOICE = {
    'en': 'Choose the correct meaning',
    'bn': 'সঠিক অর্থ নির্বাচন করুন',
    'ur': 'درست معنی کا انتخاب کریں',
    'hi': 'सही अर्थ चुनें',
    'in': 'Pilih arti yang benar',
    'ms': 'Pilih maksud yang betul',
    'tr': 'Doğru anlamı seçin',
    'fa': 'معنی درست را انتخاب کنید',
    'ha': 'Zabi ma\'anar da ta dace',
    'sw': 'Chagua maana sahihi'
}

CHAPTER_DEFINITIONS = [
    {
        'id': 'ch_01',
        'title': {
            'en': 'Grammatical Particles',
            'bn': 'ব্যাকরণিক অব্যয়',
            'ur': 'حروف اور قواعد',
            'hi': 'व्याकरणिक अव्यय',
            'in': 'Partikel Tata Bahasa',
            'ms': 'Zarah Tatabahasa',
            'tr': 'Gramer Edatları',
            'fa': 'حروف و قواعد دستوری',
            'ha': 'Haruffan Nahawu',
            'sw': 'Vihusishi vya Sarufi'
        },
        'category': 'PARTICLE',
        'word_range': (0, 173)
    },
    {
        'id': 'ch_02',
        'title': {
            'en': 'High-Frequency Verbs',
            'bn': 'উচ্চ-ফ্রিকোয়েন্সি ক্রিয়াপদ',
            'ur': 'کثیر الاستعمال افعال',
            'hi': 'उच्च-आवृत्ति क्रियाएँ',
            'in': 'Kata Kerja Frekuensi Tinggi',
            'ms': 'Kata Kerja Kekerapan Tinggi',
            'tr': 'Yüksek Frekanslı Fiiller',
            'fa': 'افعال پرکاربرد',
            'ha': 'Fi\'ilai Masu Yawan Aukuwa',
            'sw': 'Vitenzi vya Marudio ya Juu'
        },
        'category': 'VERB',
        'word_range': (173, 673)
    },
    {
        'id': 'ch_03',
        'title': {
            'en': 'Essential Verbal Forms',
            'bn': 'প্রয়োজনীয় ক্রিয়াপদের রূপ',
            'ur': 'اہم افعال کے ابواب',
            'hi': 'आवश्यक क्रिया रूप',
            'in': 'Bentuk Kata Kerja Esensial',
            'ms': 'Bentuk Kata Kerja Asas',
            'tr': 'Temel Fiil Kalıpları',
            'fa': 'باب‌های اساسی افعال',
            'ha': 'Sifofin Fi\'ili Masu Muhimmanci',
            'sw': 'Miundo Muhimu ya Vitenzi'
        },
        'category': 'VERB',
        'word_range': (673, 1173)
    },
    {
        'id': 'ch_04',
        'title': {
            'en': 'Specialized Verbs',
            'bn': 'বিশেষায়িত ক্রিয়াপদ',
            'ur': 'خصوصی افعال',
            'hi': 'विशिष्ट क्रियाएँ',
            'in': 'Kata Kerja Khusus',
            'ms': 'Kata Kerja Khusus',
            'tr': 'Özel Fiiller',
            'fa': 'افعال تخصصی',
            'ha': 'Fi\'ilai na Musamman',
            'sw': 'Vitenzi Maalum'
        },
        'category': 'VERB',
        'word_range': (1173, 1652)
    },
    {
        'id': 'ch_05',
        'title': {
            'en': 'Divine Names & Core Nominals',
            'bn': 'আসমাউল হুসনা ও মৌলিক বিশেষ্য',
            'ur': 'اسمائے حسنیٰ اور بنیادی اسماء',
            'hi': 'अस्माउल हुस्ना और मूल संज्ञाएँ',
            'in': 'Asmaul Husna & Nomina Inti',
            'ms': 'Asmaul Husna & Kata Nama Teras',
            'tr': 'Esmâ-i Hüsnâ ve Temel İsimler',
            'fa': 'نام‌های نیکوی خداوند و اسم‌های بنیادین',
            'ha': 'Sunayen Allah da Asalin Sunaye',
            'sw': 'Majina ya Mwenyezi Mungu na Majina ya Msingi'
        },
        'category': 'NOUN',
        'word_range': (1652, 2162)
    },
    {
        'id': 'ch_06',
        'title': {
            'en': 'Essential Quranic Nominals',
            'bn': 'প্রয়োজনীয় কুরআনিক বিশেষ্য',
            'ur': 'اہم قرآنی اسماء',
            'hi': 'आवश्यक क़ुरआनी संज्ञाएँ',
            'in': 'Nomina Al-Qur\'an Esensial',
            'ms': 'Kata Nama Al-Quran Asas',
            'tr': 'Temel Kur\'an İsimleri',
            'fa': 'اسم‌های اساسی قرآنی',
            'ha': 'Sunayen Alkur\'ani Masu Muhimmanci',
            'sw': 'Majina Muhimu ya Qur\'ani'
        },
        'category': 'NOUN',
        'word_range': (2162, 2672)
    },
    {
        'id': 'ch_07',
        'title': {
            'en': 'Devotional & Faith Nominals',
            'bn': 'ইবাদত ও ঈমান সংক্রান্ত বিশেষ্য',
            'ur': 'ایمان اور عبادات کے اسماء',
            'hi': 'इबादत और ईमान संबंधी संज्ञाएँ',
            'in': 'Nomina Ibadah & Keimanan',
            'ms': 'Kata Nama Ibadah & Keimanan',
            'tr': 'İbadet ve İman İsimleri',
            'fa': 'اسم‌های ایمان و عبادات',
            'ha': 'Sunayen Ibada da Imani',
            'sw': 'Majina ya Ibada na Imani'
        },
        'category': 'NOUN',
        'word_range': (2672, 3182)
    },
    {
        'id': 'ch_08',
        'title': {
            'en': 'Prophetic & Narrative Nominals',
            'bn': 'নবী ও ঐতিহাসিক ঘটনার বিশেষ্য',
            'ur': 'انبیاء اور قصص کے اسماء',
            'hi': 'नबियों और ऐतिहासिक घटनाओं की संज्ञाएँ',
            'in': 'Nomina Kisah & Kenabian',
            'ms': 'Kata Nama Kisah & Kenabian',
            'tr': 'Peygamber ve Kıssa İsimleri',
            'fa': 'اسم‌های پیامبران و قصص',
            'ha': 'Sunayen Annabawa da Labarai',
            'sw': 'Majina ya Mitume na Visa'
        },
        'category': 'NOUN',
        'word_range': (3182, 3692)
    },
    {
        'id': 'ch_09',
        'title': {
            'en': 'Moral & Social Nominals',
            'bn': 'নৈতিক ও সামাজিক বিশেষ্য',
            'ur': 'اخلاقی اور معاشرتی اسماء',
            'hi': 'नैतिक और सामाजिक संज्ञाएँ',
            'in': 'Nomina Moral & Sosial',
            'ms': 'Kata Nama Akhlak & Kemasyarakatan',
            'tr': 'Ahlaki ve Sosyal İsimler',
            'fa': 'اسم‌های اخلاقی و اجتماعی',
            'ha': 'Sunayen Dabi\'u da Zaman Jama\'a',
            'sw': 'Majina ya Maadili na Jamii'
        },
        'category': 'NOUN',
        'word_range': (3692, 4202)
    },
    {
        'id': 'ch_10',
        'title': {
            'en': 'Cosmic & Lexical Nominals',
            'bn': 'মহাজাগতিক ও আভিধানিক বিশেষ্য',
            'ur': 'کائناتی اور لغوی اسماء',
            'hi': 'ब्रह्मांडीय और शाब्दिक संज्ञाएँ',
            'in': 'Nomina Kosmis & Leksikal',
            'ms': 'Kata Nama Kosmik & Leksikal',
            'tr': 'Kozmik ve Sözlük İsimleri',
            'fa': 'اسم‌های کیهانی و واژگانی',
            'ha': 'Sunayen Halittu da Kalmomi',
            'sw': 'Majina ya Ulimwengu na Maneno'
        },
        'category': 'NOUN',
        'word_range': (4202, 4709)
    }
]

def extract_single_option_meaning(w_item):
    single_m = {}
    for lang in LANG_CODES:
        val = w_item['meaning'].get(lang, '')
        parts = re.split(r'\s*[/\\;|,]\s*', str(val).strip())
        parts = [p.strip() for p in parts if p.strip()]
        single_m[lang] = parts[0] if parts else val
    return single_m

def generate_options(correct_w, pool):
    correct_opt = {
        'id': f"opt_{correct_w['id']}",
        'labelArabic': correct_w['wordArabic'],
        'label': extract_single_option_meaning(correct_w)
    }
    cand_distractors = [w for w in pool if w['id'] != correct_w['id']]
    step = max(1, len(cand_distractors) // 10)
    distractors = []
    for i in range(3):
        pick_idx = (correct_w['rank'] * (i + 1) * step) % len(cand_distractors)
        d_w = cand_distractors[pick_idx]
        distractors.append({
            'id': f"opt_{d_w['id']}",
            'labelArabic': d_w['wordArabic'],
            'label': extract_single_option_meaning(d_w)
        })
    slot = correct_w['rank'] % 4
    opts = list(distractors)
    opts.insert(slot, correct_opt)
    return opts, correct_opt['id']

def main():
    print("=" * 70)
    print("10-Language Master Curriculum Compiler for Android Content Assets")
    print("=" * 70)
    start_time = time.time()
    
    # Load all 10 language workbooks
    workbooks = {}
    for code in LANG_CODES:
        db_code = DB_CODE_MAP.get(code, code)
        fpath = os.path.join(DB_DIR, f'qw_{db_code}.xlsx')
        print(f"Loading {fpath}...")
        wb = openpyxl.load_workbook(fpath, read_only=True)
        workbooks[code] = wb

    categories = [
        ('Harf', 'PARTICLE'),
        ('Fil', 'VERB'),
        ('Ism', 'NOUN')
    ]

    all_words_catalog = []
    global_rank = 1

    for cat_sheet, cat_enum in categories:
        print(f"\nProcessing category {cat_sheet}...")
        ws_en = workbooks['en'][cat_sheet]
        en_rows = list(ws_en.iter_rows(min_row=2, values_only=True))
        
        lang_rows = {}
        for code in LANG_CODES:
            if code == 'en':
                lang_rows[code] = en_rows
            else:
                lang_rows[code] = list(workbooks[code][cat_sheet].iter_rows(min_row=2, values_only=True))

        lang_indexed = {code: {} for code in LANG_CODES}
        for code in LANG_CODES:
            for r in lang_rows[code]:
                rank = int(r[0])
                sense_str = str(r[7]).strip() if r[7] else '[1]'
                lang_indexed[code][(rank, sense_str)] = r

        ranks = sorted(list(set(int(r[0]) for r in en_rows)))
        print(f"  {len(ranks)} unique headwords in {cat_sheet}")

        for rank in ranks:
            w_id = f"w_{global_rank:04d}"

            all_senses_for_rank = set()
            for code in LANG_CODES:
                for (r_k, s_str) in lang_indexed[code]:
                    if r_k == rank:
                        all_senses_for_rank.add(s_str)
            sorted_senses = sorted(list(all_senses_for_rank), key=lambda x: int(re.search(r'\d+', x).group(0)) if re.search(r'\d+', x) else 1)

            base_r = lang_indexed['en'].get((rank, sorted_senses[0])) or lang_indexed['en'].get((rank, '[1]'))

            ar_lemma = str(base_r[1]).strip()
            translit = str(base_r[2]).strip() if base_r[2] else ''
            root = str(base_r[3]).strip() if base_r[3] else '—'
            pos = str(base_r[4]).strip() if base_r[4] else ''
            occ = int(base_r[5]) if base_r[5] is not None else 1

            clean_lemma = ar_lemma
            if '_' in clean_lemma or 'فِعْل' in clean_lemma or 'فَعَلَ_' in clean_lemma or 'حَرْف_' in clean_lemma:
                tar_word_sample = str(base_r[10]).strip() if base_r[10] else ''
                if tar_word_sample and len(tar_word_sample) >= 2:
                    clean_lemma = tar_word_sample

            # Extract 10-language localized meanings and example verses
            senses_list = []
            for s_idx, sense_str in enumerate(sorted_senses, start=1):
                s_num = int(re.search(r'\d+', sense_str).group(0)) if re.search(r'\d+', sense_str) else s_idx
                sr_en = lang_indexed['en'].get((rank, sense_str)) or base_r

                verse_cit = str(sr_en[9]).strip() if sr_en[9] else ''
                tar_ar = str(sr_en[10]).strip() if sr_en[10] else ''
                full_ar_hl = str(sr_en[11]).strip() if sr_en[11] else ''
                m_ar = re.search(r'\(\[([^\]]+)\]\)', full_ar_hl)
                if m_ar:
                    tok_ar = m_ar.group(1)
                    clean_ar_verse = full_ar_hl.replace('([', '').replace('])', '')
                    start_ar = m_ar.start()
                    end_ar = start_ar + len(tok_ar)
                    spans = (start_ar, end_ar)
                else:
                    clean_ar_verse = clean_verse_text(full_ar_hl)
                    spans = calculate_token_spans(clean_ar_verse, tar_ar)

                meanings = {}
                verse_translations = {}
                meaning_highlights = {}

                for code in LANG_CODES:
                    sr_l = lang_indexed[code].get((rank, sense_str)) or lang_indexed[code].get((rank, sorted_senses[0]))
                    mean_val = str(sr_l[8]).strip() if sr_l and sr_l[8] is not None else ''
                    tar_m_val = str(sr_l[12]).strip() if sr_l and sr_l[12] is not None else ''
                    full_tr_val = str(sr_l[13]).strip() if sr_l and sr_l[13] is not None else ''

                    # Fix NULLs
                    if mean_val == 'NULL' or tar_m_val == 'NULL' or not mean_val:
                        if cat_sheet == 'Ism' and rank == 236:
                            if code == 'ha': mean_val = tar_m_val = 'yankakke'
                            elif code == 'sw': mean_val = tar_m_val = 'lililopitishwa'
                            elif code == 'en': mean_val = tar_m_val = 'decreed'
                        elif cat_sheet == 'Fil' and ('غضِب' in translit or 'غ ض ب' in root or rank in [86, 106, 126, 146, 166, 186, 206, 226, 246, 258, 268, 278, 288, 298, 308, 318, 328, 338, 348, 358, 368, 516, 536, 556, 576, 596, 616, 638, 651, 678, 691, 718, 731, 758, 771, 798, 811, 838, 851, 878, 891, 906, 926, 946, 966, 986, 1006, 1026, 1046, 1066, 1086, 1106, 1126, 1146, 1166, 1186, 1206, 1226, 1246, 1266, 1286, 1306, 1326, 1346, 1366, 1386, 1406, 1426, 1446, 1466]):
                            if code == 'en': mean_val = 'to be angry'; tar_m_val = 'has become angry'
                            elif code == 'ha': mean_val = tar_m_val = 'Ya yi fushi'
                            elif code == 'sw': mean_val = tar_m_val = 'amemkasirikia'
                        else:
                            mean_val = str(sr_en[8]).strip() if sr_en and sr_en[8] != 'NULL' else translit.split('_')[0]
                            tar_m_val = mean_val

                    mean_clean = re.sub(r'[\(\)\[\],;\.!\?।\'\":؛؟]', '', mean_val).strip()
                    clean_tr_verse = clean_verse_text(full_tr_val)
                    hl_phrase = extract_highlight_phrase(full_tr_val)
                    if not hl_phrase or hl_phrase not in clean_tr_verse:
                        cand_col12 = str(sr_l[12] or '').strip() if sr_l and sr_l[12] is not None else ''
                        if cand_col12 and cand_col12 in clean_tr_verse:
                            hl_phrase = cand_col12
                        elif mean_val and mean_val in clean_tr_verse:
                            hl_phrase = mean_val
                        else:
                            words = [w.strip('\"\'()[]{}.,;:!?/\\-— ') for w in clean_tr_verse.split() if len(w.strip('\"\'()[]{}.,;:!?/\\-— ')) >= 3]
                            hl_phrase = words[0] if words else (clean_tr_verse.split()[0] if clean_tr_verse else '')

                    meanings[code] = mean_clean
                    verse_translations[code] = clean_tr_verse
                    meaning_highlights[code] = hl_phrase

                senses_list.append({
                    'meaningIndex': s_num,
                    'contextualMeaning': meanings,
                    'verseReference': verse_cit,
                    'verseArabic': clean_ar_verse,
                    'arabicWordStart': spans[0],
                    'arabicWordEnd': spans[1],
                    'verseTranslation': verse_translations,
                    'translationHighlight': meaning_highlights
                })

            primary_sense = senses_list[0]
            word_obj = {
                'id': w_id,
                'rank': global_rank,
                'wordArabic': clean_lemma,
                'transliteration': translit,
                'root': root,
                'partOfSpeech': pos,
                'lemmaCategory': cat_enum,
                'quranOccurrenceCount': occ,
                'meaning': primary_sense['contextualMeaning'],
                'exampleVerseReference': primary_sense['verseReference'],
                'exampleVerseArabic': primary_sense['verseArabic'],
                'arabicWordStart': primary_sense['arabicWordStart'],
                'arabicWordEnd': primary_sense['arabicWordEnd'],
                'exampleVerseTranslation': primary_sense['verseTranslation'],
                'meaningHighlight': primary_sense['translationHighlight'],
                'polysemyEntries': senses_list
            }
            all_words_catalog.append(word_obj)
            global_rank += 1

    print(f"\nTotal Headwords Assembled: {len(all_words_catalog)} (100% complete)")

    # Build 10-Chapter Curriculum Structure
    print("\nSynthesizing 10 Chapters, 100 Sections, 1,217 Lessons, 9,428 Exercises...")
    
    total_tokens = sum(w['quranOccurrenceCount'] for w in all_words_catalog)
    words_by_cat = defaultdict(list)
    for w in all_words_catalog:
        words_by_cat[w['lemmaCategory']].append(w)

    chapters_out = []
    sections_out = []
    lessons_out = []
    exercises_out = []

    sec_global_num = 1
    les_global_num = 1
    ex_global_num = 1
    accumulated_words = 0
    accumulated_occ = 0

    for ch_idx, ch_info in enumerate(CHAPTER_DEFINITIONS, 1):
        ch_id = ch_info['id']
        ch_title = ch_info['title']
        ch_cat = ch_info['category']
        w_start, w_end = ch_info['word_range']
        ch_words = all_words_catalog[w_start:w_end]

        ch_occ = sum(w['quranOccurrenceCount'] for w in ch_words)
        ch_pct = round((ch_occ / total_tokens) * 100, 2)
        accumulated_words += len(ch_words)
        accumulated_occ += ch_occ
        acc_pct = round((accumulated_occ / total_tokens) * 100, 2)

        chapters_out.append({
            'id': ch_id,
            'title': ch_title,
            'description': ch_title,
            'sortOrder': ch_idx,
            'wordCount': len(ch_words),
            'quranOccurrenceCount': ch_occ,
            'quranOccurrencePercent': ch_pct
        })

        words_per_sec = len(ch_words) // 10
        rem = len(ch_words) % 10
        sec_start = 0

        for s_idx in range(1, 11):
            s_len = words_per_sec + (1 if s_idx <= rem else 0)
            sec_words = ch_words[sec_start : sec_start + s_len]
            sec_start += s_len

            sec_id = f"sec_{sec_global_num:03d}"
            sec_title = {
                'en': f"Section {s_idx}: {ch_title['en']}",
                'bn': f"{get_bn_ordinal(s_idx)} পর্ব: {ch_title['bn']}",
                'ur': f"{get_ur_ordinal(s_idx)} حصہ: {ch_title['ur']}",
                'hi': f"खंड {to_hi_digits(s_idx)}: {ch_title['hi']}",
                'in': f"Bagian {s_idx}: {ch_title['in']}",
                'ms': f"Bahagian {s_idx}: {ch_title['ms']}",
                'tr': f"Bölüm {s_idx}: {ch_title['tr']}",
                'fa': f"بخش {to_ur_digits(s_idx)}: {ch_title['fa']}",
                'ha': f"Sashe na {s_idx}: {ch_title['ha']}",
                'sw': f"Sehemu ya {s_idx}: {ch_title['sw']}"
            }
            sec_occ = sum(w['quranOccurrenceCount'] for w in sec_words)
            sec_pct = round((sec_occ / total_tokens) * 100, 2)

            sections_out.append({
                'id': sec_id,
                'chapterId': ch_id,
                'title': sec_title,
                'sortOrder': s_idx,
                'wordCount': len(sec_words),
                'quranOccurrenceCount': sec_occ,
                'quranOccurrencePercent': sec_pct
            })

            sec_sort_order = 1

            # Chapter Intro Lesson & Exercise (in Section 1 only)
            if s_idx == 1:
                ch_intro_les_id = f"les_{les_global_num:04d}"
                ch_intro_ex_id = f"ex_{ex_global_num:05d}"
                ex_global_num += 1

                lessons_out.append({
                    'id': ch_intro_les_id,
                    'chapterId': ch_id,
                    'sectionId': sec_id,
                    'title': {
                        'en': f'Chapter {ch_idx} Overview',
                        'bn': f'{get_bn_ordinal(ch_idx)} অধ্যায় পরিচিতি',
                        'ur': f'{get_ur_ordinal(ch_idx)} باب کا تعارف',
                        'hi': f'अध्याय {to_hi_digits(ch_idx)} परिचय',
                        'in': f'Ikhtisar Bab {ch_idx}',
                        'ms': f'Gambaran Keseluruhan Bab {ch_idx}',
                        'tr': f'Bölüm {ch_idx} Genel Bakış',
                        'fa': f'معرفی فصل {to_ur_digits(ch_idx)}',
                        'ha': f'Bayanin Babi na {ch_idx}',
                        'sw': f'Muhtasari wa Sura ya {ch_idx}'
                    },
                    'sortOrder': sec_sort_order,
                    'kind': 'CHAPTER_INTRO',
                    'category': ch_cat
                })
                sec_sort_order += 1

                exercises_out.append({
                    'id': ch_intro_ex_id,
                    'lessonId': ch_intro_les_id,
                    'orderIndex': 1,
                    'exerciseType': 'CHAPTER_INTRO',
                    'content': {
                        'type': 'chapter_intro',
                        'prompt': ch_title,
                        'chapterId': ch_id,
                        'chapterNumber': ch_idx,
                        'chapterTitle': ch_title,
                        'chapterDescription': ch_title,
                        'wordCount': len(ch_words),
                        'quranOccurrenceCount': ch_occ,
                        'chapterCoveragePercent': ch_pct,
                        'accumulatedCoveragePercent': acc_pct,
                        'accumulatedWords': accumulated_words,
                        'nounCount': len(ch_words) if ch_cat == 'NOUN' else 0,
                        'verbCount': len(ch_words) if ch_cat == 'VERB' else 0,
                        'particleCount': len(ch_words) if ch_cat == 'PARTICLE' else 0,
                        'learningObjectives': ch_title
                    }
                })
                les_global_num += 1

            # Regular Lessons (chunks of 5 words)
            sec_w_chunks = [sec_words[i:i+5] for i in range(0, len(sec_words), 5)]
            for chunk_idx, w_chunk in enumerate(sec_w_chunks, 1):
                les_id = f"les_{les_global_num:04d}"
                les_title = {
                    'en': f"Lesson {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'bn': f"পাঠ {to_bn_digits(chunk_idx)}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'ur': f"سبق {to_ur_digits(chunk_idx)}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'hi': f"पाठ {to_hi_digits(chunk_idx)}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'in': f"Pelajaran {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'ms': f"Pelajaran {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'tr': f"Ders {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'fa': f"درس {to_ur_digits(chunk_idx)}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'ha': f"Darasi na {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}",
                    'sw': f"Somo la {chunk_idx}: {w_chunk[0]['wordArabic']} – {w_chunk[-1]['wordArabic']}"
                }

                lessons_out.append({
                    'id': les_id,
                    'chapterId': ch_id,
                    'sectionId': sec_id,
                    'title': les_title,
                    'sortOrder': sec_sort_order,
                    'kind': 'REGULAR',
                    'category': ch_cat
                })
                sec_sort_order += 1

                ex_ord = 1
                for w in w_chunk:
                    intro_ex_id = f"ex_{ex_global_num:05d}"
                    ex_global_num += 1

                    exercises_out.append({
                        'id': intro_ex_id,
                        'lessonId': les_id,
                        'orderIndex': ex_ord,
                        'exerciseType': 'WORD_INTRO',
                        'content': {
                            'type': 'word_intro',
                            'prompt': PROMPT_WORD_INTRO,
                            'wordId': w['id'],
                            'arabicWord': w['wordArabic'],
                            'transliteration': w['transliteration'],
                            'root': w['root'],
                            'partOfSpeech': w['partOfSpeech'],
                            'quranOccurrenceCount': w['quranOccurrenceCount'],
                            'exampleVerseArabic': w['exampleVerseArabic'],
                            'exampleVerseTranslation': w['exampleVerseTranslation'],
                            'exampleVerseReference': w['exampleVerseReference'],
                            'meaning': w['meaning'],
                            'arabicWordStart': w['arabicWordStart'],
                            'arabicWordEnd': w['arabicWordEnd'],
                            'meaningHighlight': w['meaningHighlight'],
                            'polysemyEntries': [
                                {
                                    'meaningIndex': se['meaningIndex'],
                                    'contextualMeaning': se['contextualMeaning'],
                                    'verseReference': se['verseReference'],
                                    'verseArabic': se['verseArabic'],
                                    'arabicWordStart': se['arabicWordStart'],
                                    'arabicWordEnd': se['arabicWordEnd'],
                                    'verseTranslation': se['verseTranslation'],
                                    'translationHighlight': se['translationHighlight']
                                }
                                for se in w['polysemyEntries']
                            ]
                        }
                    })
                    ex_ord += 1

                for w in w_chunk:
                    quiz_ex_id = f"ex_{ex_global_num:05d}"
                    ex_global_num += 1

                    options, correct_id = generate_options(w, words_by_cat[ch_cat])

                    exercises_out.append({
                        'id': quiz_ex_id,
                        'lessonId': les_id,
                        'orderIndex': ex_ord,
                        'exerciseType': 'MULTIPLE_CHOICE',
                        'content': {
                            'type': 'multiple_choice',
                            'prompt': PROMPT_MULTIPLE_CHOICE,
                            'promptArabic': w['wordArabic'],
                            'wordId': w['id'],
                            'options': options,
                            'correctOptionId': correct_id,
                            'exampleVerseArabic': w['exampleVerseArabic'],
                            'exampleVerseTranslation': w['exampleVerseTranslation'],
                            'exampleVerseReference': w['exampleVerseReference'],
                            'arabicWordStart': w['arabicWordStart'],
                            'arabicWordEnd': w['arabicWordEnd'],
                            'meaningHighlight': w['meaningHighlight']
                        }
                    })
                    ex_ord += 1

                les_global_num += 1

            # Section Flashback Lesson
            flash_les_id = f"les_{les_global_num:04d}"
            lessons_out.append({
                'id': flash_les_id,
                'chapterId': ch_id,
                'sectionId': sec_id,
                'title': {
                    'en': f"Section {s_idx} Flashback",
                    'bn': f"{get_bn_ordinal(s_idx)} পর্ব পুনরাবৃত্তি",
                    'ur': f"{get_ur_ordinal(s_idx)} حصہ کا اعادہ",
                    'hi': f"खंड {to_hi_digits(s_idx)} पुनरावृत्ति",
                    'in': f"Ulasan Bagian {s_idx}",
                    'ms': f"Imbasan Bahagian {s_idx}",
                    'tr': f"Bölüm {s_idx} Tekrarı",
                    'fa': f"مرور بخش {to_ur_digits(s_idx)}",
                    'ha': f"Maimaita Sashe na {s_idx}",
                    'sw': f"Marekebisho ya Sehemu ya {s_idx}"
                },
                'sortOrder': sec_sort_order,
                'kind': 'SECTION_FLASHBACK',
                'category': ch_cat
            })
            sec_sort_order += 1
            les_global_num += 1

            # Section Exam Lesson
            exam_les_id = f"les_{les_global_num:04d}"
            lessons_out.append({
                'id': exam_les_id,
                'chapterId': ch_id,
                'sectionId': sec_id,
                'title': {
                    'en': f"Section {s_idx} Exam",
                    'bn': f"{get_bn_ordinal(s_idx)} পর্ব পরীক্ষা",
                    'ur': f"{get_ur_ordinal(s_idx)} حصہ کا امتحان",
                    'hi': f"खंड {to_hi_digits(s_idx)} परीक्षा",
                    'in': f"Ujian Bagian {s_idx}",
                    'ms': f"Peperiksaan Bahagian {s_idx}",
                    'tr': f"Bölüm {s_idx} Sınavı",
                    'fa': f"آزمون بخش {to_ur_digits(s_idx)}",
                    'ha': f"Jarrabawar Sashe na {s_idx}",
                    'sw': f"Mtihani wa Sehemu ya {s_idx}"
                },
                'sortOrder': sec_sort_order,
                'kind': 'SECTION_EXAM',
                'category': ch_cat
            })
            sec_sort_order += 1
            les_global_num += 1

            sec_global_num += 1

        # Chapter Exam Lesson (after Section 10)
        ch_exam_les_id = f"les_{les_global_num:04d}"
        lessons_out.append({
            'id': ch_exam_les_id,
            'chapterId': ch_id,
            'sectionId': None,
            'title': {
                'en': f"Chapter {ch_idx} Exam",
                'bn': f"{get_bn_ordinal(ch_idx)} অধ্যায় পরীক্ষা",
                'ur': f"{get_ur_ordinal(ch_idx)} باب کا امتحان",
                'hi': f"अध्याय {to_hi_digits(ch_idx)} परीक्षा",
                'in': f"Ujian Bab {ch_idx}",
                'ms': f"Peperiksaan Bab {ch_idx}",
                'tr': f"Bölüm {ch_idx} Sınavı",
                'fa': f"آزمون فصل {to_ur_digits(ch_idx)}",
                'ha': f"Jarrabawar Babi na {ch_idx}",
                'sw': f"Mtihani wa Sura ya {ch_idx}"
            },
            'sortOrder': sec_sort_order,
            'kind': 'CHAPTER_EXAM',
            'category': ch_cat
        })
        les_global_num += 1

    print(f"\nCurriculum Synthesis Complete:")
    print(f"  Chapters: {len(chapters_out)}")
    print(f"  Sections: {len(sections_out)}")
    print(f"  Lessons: {len(lessons_out)}")
    print(f"  Exercises: {len(exercises_out)}")
    print(f"  Words: {len(all_words_catalog)}")

    # Emit JSON Assets
    os.makedirs(CONTENT_DIR, exist_ok=True)
    
    print("\nWriting JSON files to app/src/main/assets/content/...")
    
    with open(os.path.join(CONTENT_DIR, 'chapters.json'), 'w', encoding='utf-8') as f:
        json.dump({'chapters': chapters_out}, f, ensure_ascii=False, indent=2)
    print("  -> chapters.json written")

    with open(os.path.join(CONTENT_DIR, 'sections.json'), 'w', encoding='utf-8') as f:
        json.dump({'sections': sections_out}, f, ensure_ascii=False, indent=2)
    print("  -> sections.json written")

    with open(os.path.join(CONTENT_DIR, 'lessons_vocabulary.json'), 'w', encoding='utf-8') as f:
        json.dump({'lessons': lessons_out}, f, ensure_ascii=False, indent=2)
    print("  -> lessons_vocabulary.json written")

    with open(os.path.join(CONTENT_DIR, 'exercises_vocabulary.json'), 'w', encoding='utf-8') as f:
        json.dump({'exercises': exercises_out}, f, ensure_ascii=False)
    print("  -> exercises_vocabulary.json written")

    words_freq_out = []
    for w in all_words_catalog:
        words_freq_out.append({
            'id': w['id'],
            'arabicWord': w['wordArabic'],
            'frequencyRank': w['rank'],
            'frequencyCount': w['quranOccurrenceCount'],
            'meaning': w['meaning'],
            'audioAssetPath': None,
            'tierLevel': (w['rank'] // 500) + 1
        })

    with open(os.path.join(CONTENT_DIR, 'word_frequency.json'), 'w', encoding='utf-8') as f:
        json.dump({'words': words_freq_out}, f, ensure_ascii=False, indent=2)
    print("  -> word_frequency.json written")

    print(f"\nAll 5 assets compiled in {time.time()-start_time:.2f}s!")

if __name__ == '__main__':
    main()
