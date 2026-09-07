import openpyxl, re

wb = openpyxl.load_workbook('db/qw_en.xlsx')
ws = wb['Ism']

# 1. First restore root letters mistakenly stripped in Ism:
# 2:47: تَضَّلْتُكُمْ / ضَّلْتُكُمْ -> فَضَّلْتُكُمْ
# 2:76: تَحَ -> فَتَحَ
for r in range(2, ws.max_row + 1):
    cit = str(ws.cell(r, 10).value or '').strip()
    t_ar = str(ws.cell(r, 11).value or '').strip()
    v_ar = str(ws.cell(r, 12).value or '')
    
    if cit == 'Al-Baqarah 2:47' and 'ضَّلْتُكُمْ' in t_ar:
        ws.cell(r, 11).value = 'فَضَّلْتُكُمْ'
        ws.cell(r, 12).value = v_ar.replace('([ضَّلْتُكُمْ])', '([فَضَّلْتُكُمْ])').replace('فَ([ضَّلْتُكُمْ])', '([فَضَّلْتُكُمْ])')
    elif cit == 'Al-Baqarah 2:76' and t_ar == 'تَحَ':
        ws.cell(r, 11).value = 'فَتَحَ'
        ws.cell(r, 12).value = v_ar.replace('([تَحَ])', '([فَتَحَ])').replace('فَ([تَحَ])', '([فَتَحَ])')

# 2. Curation dictionary for stop words in Ism
ISM_FIXES = {
    ('Al-Baqarah 2:205', 'ٱلْأَرْضِ'): ('land', r'throughout (\[the\]) land', r'throughout the ([land])'),
    ('Al-Baqarah 2:87', 'بِٱلرُّسُلِ'): ('messengers', r'after him (\[with\]) messengers', r'after him with ([messengers])'),
    ('Al-Baqarah 2:259', 'يَوْمًا'): ('a day', r'I have remained (\[a day\])|remained (\[for\]) a hundred years.*?"I have remained a day', r'I have remained ([a day])'),
    ('Al-Baqarah 2:7', 'عَذَابٌ'): ('great punishment', r'is a (\[great punishment\])|over their vision (\[is\]) a veil\. And for them is a great punishment', r'over their vision is a veil. And for them is a ([great punishment])'),
    ("Ar-Ra'd 13:6", 'لِّلنَّاسِ'): ('the people', r'forgiveness (\[for\]) the people', r'forgiveness for ([the people])'),
    ('Al-Baqarah 2:8', 'بِمُؤْمِنِينَ'): ('believers', r'they (\[are\]) not believers', r'they are not ([believers])'),
    ("Ash-Shu'ara 26:89", 'بِقَلْبٍ'): ('heart', r'with a sound (\[heart\])|comes to Allāh (\[with\]) a sound heart', r'comes to Allāh with a sound ([heart])'),
    ('Al-Baqarah 2:86', 'ٱلدُّنْيَا'): ('this world', r'the life (\[of\]) this world', r'the life of ([this world])'),
    ('Al-Baqarah 2:153', 'ٱلصَّـٰبِرِينَ'): ('the patient', r'Allāh is with (\[the\]) patient', r'Allāh is with ([the patient])'),
    ('Al-Baqarah 2:202', 'ٱلْحِسَابِ'): ('account', r'Allāh is swift (\[in\]) account', r'Allāh is swift in ([account])'),
    ('Al-Baqarah 2:243', 'أَلَمْ'): {
        'target_ar': 'أَحْيَـٰهُمْ',
        'v_ar': 'أَلَمْ تَرَ إِلَى ٱلَّذِينَ خَرَجُوا۟ مِن دِيَـٰرِهِمْ وَهُمْ أُلُوفٌ حَذَرَ ٱلْمَوْتِ فَقَالَ لَهُمُ ٱللَّهُ مُوتُوا۟ ثُمَّ ([أَحْيَـٰهُمْ]) ۚ إِنَّ ٱللَّهَ لَذُو فَضْلٍ عَلَى ٱلنَّاسِ وَلَـٰكِنَّ أَكْثَرَ ٱلنَّاسِ لَا يَشْكُرُونَ',
        'target_en': 'life',
        'sub': (r'restored them to (\[life\])|fearing death\? Allāh said to them, "Die"; then He restored them to life',
                r'restored them to ([life])')
    },
    ('Al-Baqarah 2:102', 'مُلْكِ'): ('reign', r'during (\[the\]) reign of Solomon', r'during the ([reign]) of Solomon'),
    ('Al-Baqarah 2:5', 'ٱلْمُفْلِحُونَ'): ('the successful', r'who (\[are\]) the successful', r'who are ([the successful])'),
    ('Al-Isra 17:93', 'سُبْحَانَ'): ('Exalted', r'\(\[Exalted\]\) is my Lord|Say, "(\[is\]) my Lord|Say, "Exalted (\[is\]) my Lord', r'Say, "([Exalted]) is my Lord'),
    ('Al-Baqarah 2:3', 'مِمَّا'): {
        'target_ar': 'يُؤْمِنُونَ',
        'v_ar': 'ٱلَّذِينَ ([يُؤْمِنُونَ]) بِٱلْغَيْبِ وَيُقِيمُونَ ٱلصَّلَوٰةَ وَمِمَّا رَزَقْنَـٰهُمْ يُنفِقُونَ',
        'target_en': 'believe',
        'sub': (r'Who (\[believe\]) in the unseen|Who believe in the unseen, establish prayer, (\[and\]) spend',
                r'Who ([believe]) in the unseen')
    },
    ('An-Nisa 4:168', 'طَرِيقًا'): ('a path', r'guide them (\[to\]) a path', r'guide them to ([a path])'),
    ('Al-Baqarah 2:175', 'بِٱلْمَغْفِرَةِ'): ('forgiveness', r'and (\[forgiveness\]) for punishment|error and (\[forgiveness\]) for|error and forgiveness (\[for\]) punishment', r'error and ([forgiveness]) for punishment'),
    ('An-Nisa 4:13', 'ٱلْفَوْزُ'): ('attainment', r'the great (\[attainment\])|that (\[is\]) the great attainment', r'that is the great ([attainment])'),
    ('Al-Baqarah 2:129', 'ٱلْحِكْمَةَ'): ('wisdom', r'the Book and (\[wisdom\])|the Book (\[and\]) wisdom', r'the Book and ([wisdom])'),
    ('An-Nahl 16:102', 'رُوحُ'): ('Spirit', r'The Pure (\[Spirit\])|\(\[The\]\) Pure Spirit', r'The Pure ([Spirit])'),
    ('Al-Baqarah 2:51', 'مُوسَىٰٓ'): ('Moses', r'appointment with (\[Moses\])|appointment with Moses (\[for\]) forty nights', r'appointment with ([Moses]) for forty nights'),
    ("Ali 'Imran 3:33", 'إِبْرَٰهِيمَ'): ('Abraham', r'family of (\[Abraham\])|family (\[of\]) Abraham', r'family of ([Abraham])'),
    ("Ali 'Imran 3:33", 'نُوحًا'): ('Noah', r'Adam and (\[Noah\])|Adam (\[and\]) Noah', r'Adam and ([Noah])'),
    ("Ali 'Imran 3:84", 'عِيسَىٰ'): ('Jesus', r'and (\[Jesus\]) and|revealed to Abraham.*?(?:and \(\[and\]\) Jesus|and (\[and\]) Jesus)', r'and ([Jesus]) and'),
    ("Al-An'am 6:84", 'يُوسُفَ'): ('Joseph', r'and (\[Joseph\])|and \(\[and\]\) Joseph', r'and ([Joseph])'),
    ('An-Nisa 4:163', 'سُلَيْمَـٰنَ'): ('Solomon', r'and (\[Solomon\])|and \(\[and\]\) Solomon', r'and ([Solomon])'),
    ('Al-Baqarah 2:251', 'دَاوُۥدُ'): ('David', r'and (\[David\]) killed Goliath|and David killed Goliath.*?\(\[they\]\)', r'and ([David]) killed Goliath'),
    ("Ali 'Imran 3:36", 'مَرْيَمَ'): ('Mary', r'named her (\[Mary\])|\(\[But\]\) when she delivered her', r'named her ([Mary])'),
    ("Ali 'Imran 3:96", 'بِبَكَّةَ'): ('Bakkah', r'that at (\[Bakkah\])|that (\[at\]) Bakkah', r'that at ([Bakkah])'),
    ("Ali 'Imran 3:3", 'ٱلْإِنجِيلَ'): ('Gospel', r'the Torah and the (\[Gospel\])|Torah (\[And\]) the Gospel', r'the Torah and the ([Gospel])'),
    ('An-Nisa 4:163', 'زَبُورًا'): ('Psalms', r'to David the (\[Psalms\])|David (\[the\]) Psalms', r'to David the ([Psalms])'),
    ('An-Nisa 4:82', 'ٱلْقُرْءَانَ'): ('Qur’ān', r'upon the (\[Qur’ān\])|upon (\[the\]) Qur’ān', r'upon the ([Qur’ān])'),
    ('Al-Baqarah 2:49', 'فِرْعَوْنَ'): ('Pharaoh', r'people of (\[Pharaoh\])|people (\[of\]) Pharaoh', r'people of ([Pharaoh])'),
    ('Al-Baqarah 2:97', 'لِّجِبْرِيلَ'): ('Gabriel', r'enemy to (\[Gabriel\])|enemy (\[to\]) Gabriel', r'enemy to ([Gabriel])'),
    ('Al-Baqarah 2:286', 'ٱرْحَمْنَآ'): ('have mercy upon us', r'and (\[have mercy upon us\])|and forgive us; (\[and\]) have mercy upon us', r'and forgive us; and ([have mercy upon us])'),
    ('Al-Fatihah 1:1', 'ٱلرَّحِيمِ'): ('Especially Merciful', r'the (\[Especially Merciful\])|the Entirely Merciful, (\[the\]) Especially Merciful', r'the Entirely Merciful, the ([Especially Merciful])'),
    ('Al-Baqarah 2:29', 'عَلِيمٌ'): ('Knowing', r'all things (\[Knowing\])|earth\. Then He directed.*?\(\[is\]\)', r'all things ([Knowing])'),
    ('Al-Baqarah 2:32', 'ٱلْحِكِيمُ'): ('Wise', r'the Knowing, the (\[Wise\])|the Knowing, (\[the\]) Wise', r'the Knowing, the ([Wise])'),
    ('Al-Baqarah 2:173', 'غَفُورٌ'): ('Forgiving', r'Allāh is (\[Forgiving\])|Allāh (\[is\]) Forgiving', r'Allāh is ([Forgiving])'),
    ('Al-Baqarah 2:234', 'خَبِيرٌ'): ('Acquainted', r'what you do is (\[Acquainted\])|what you do (\[is\]) Acquainted', r'what you do is ([Acquainted])'),
    ('Hud 11:73', 'مَّجِيدٌ'): ('Glorious', r'Praiseworthy and (\[Glorious\])|\(\[They\]\) said, "Are you amazed', r'Praiseworthy and ([Glorious])'),
    ('Al-Baqarah 2:129', 'ٱلْعَزِيزُ'): ('Exalted in Might', r'You are the (\[Exalted in Might\])|You are (\[the\]) Exalted in Might', r'You are the ([Exalted in Might])'),
    ('Al-Baqarah 2:35', 'تَقْرَبَا'): ('approach', r'do not (\[approach\]) this tree|dwell, (\[you\]) and your wife', r'do not ([approach]) this tree'),
    ('Al-Baqarah 2:10', 'يَكْذِبُونَ'): ('lie', r'used to (\[lie\])|because (\[they\]) \[habitually\] used to lie', r'used to ([lie])'),
    ('Al-Baqarah 2:105', 'ٱلْمُشْرِكِينَ'): ('polytheists', r'or the (\[polytheists\])|Neither (\[those\]) who disbelieve', r'or the ([polytheists])'),
    ('Al-Baqarah 2:267', 'ٱلْخَبِيثَ'): ('bad', r'aim toward the (\[bad\])|aim toward (\[the\]) bad', r'aim toward the ([bad])'),
    ('Al-Baqarah 2:180', 'ٱلْأَقْرَبِينَ'): ('near relatives', r'parents and (\[near relatives\])|parents (\[and\]) near relatives', r'parents and ([near relatives])'),
    ('Al-Baqarah 2:20', 'أَظْلَمَ'): ('darkens', r'when it (\[darkens\]) to them|Every time (\[it\]) lights', r'when it ([darkens]) to them'),
    ("Ali 'Imran 3:7", 'مُّحْكَمَـٰتٌ'): ('precise', r'verses \[that are\] (\[precise\])|has sent down to (\[you\])', r'verses [that are] ([precise])'),
    ('Al-Baqarah 2:58', 'نَّغْفِرْ'): ('forgive', r'We will (\[forgive\]) you|when (\[We\]) said, "Enter', r'We will ([forgive]) you'),
    ('Al-Baqarah 2:37', 'ٱلتَّوَّابُ'): ('Accepting of repentance', r'is the (\[Accepting of repentance\])|Indeed, it (\[is\]) He who is the Accepting of repentance', r'Indeed, it is He who is the ([Accepting of repentance])'),
    ('Al-Baqarah 2:52', 'عَفَوْنَا'): ('forgave', r'Then We (\[forgave\]) you|Then (\[We\]) forgave you', r'Then We ([forgave]) you'),
    ('Al-Anfal 8:3', 'ٱلَّذِينَ'): {
        'target_ar': 'رَزَقْنَـٰهُمْ',
        'v_ar': 'ٱلَّذِينَ يُقِيمُونَ ٱلصَّلَوٰةَ وَمِمَّا ([رَزَقْنَـٰهُمْ]) يُنفِقُونَ',
        'target_en': 'have provided them',
        'sub': (r'We (\[have provided them\])|The ones (\[who\]) establish prayer, and from what We have provided them',
                r'The ones who establish prayer, and from what We ([have provided them]), they spend.')
    },
    ("Ali 'Imran 3:144", 'مَا'): {
        'target_ar': 'عَقِبَيْهِ',
        'v_ar': 'وَمَا مُحَمَّدٌ إِلَّا رَسُولٌ قَدْ خَلَتْ مِن قَبْلِهِ ٱلرُّسُلُ ۚ أَفَإِي۟ن مَّاتَ أَوْ قُتِلَ ٱنقَلَبْتُمْ عَلَىٰٓ أَعْقَـٰبِكُمْ ۚ وَمَن يَنقَلِبْ عَلَىٰ ([عَقِبَيْهِ]) فَلَن يَضُرَّ ٱللَّهَ شَيْـًٔا ۗ وَسَيَجْزِى ٱللَّهُ ٱلشَّـٰكِرِينَ',
        'target_en': 'his heels',
        'sub': (r'back on (\[his heels\])|Muḥammad is not but a messenger.*?(?:(\[And\]) he who turns back|he who turns back on his heels)',
                r'he who turns back on ([his heels])')
    },
    ('Al-Baqarah 2:117', 'بَدِيعُ'): ('Originator', r'(\[Originator\]) of the heavens|Originator of (\[the\]) heavens', r'([Originator]) of the heavens'),
    ("Ali 'Imran 3:6", 'يُصَوِّرُكُمْ'): ('forms you', r'He who (\[forms you\]) in the wombs|He who forms (\[you\]) in the wombs', r'He who ([forms you]) in the wombs'),
    ("Ali 'Imran 3:44", 'نُوحِيهِ'): ('reveal', r'which We (\[reveal\]) to you|unseen which (\[We\]) reveal to you', r'unseen which We ([reveal]) to you'),
    ('Al-Fatihah 1:7', 'ٱلْمَغْضُوبِ'): ('earned [Your] anger', r'have (\[earned \[Your\] anger\])|not of (\[those\]) who have earned \[Your\] anger', r'not of those who have ([earned [Your] anger])'),
    ('Al-Baqarah 2:58', 'ٱلْمُحْسِنِينَ'): ('doers of good', r'increase the (\[doers of good\])|increase (\[the\]) doers of good', r'increase the ([doers of good])'),
    ('Al-Baqarah 2:49', 'يَسُومُونَكُمْ'): ('afflicted you', r'who (\[afflicted you\]) with|Pharaoh, (\[who\]) afflicted you with', r'Pharaoh, who ([afflicted you]) with'),
    ('Al-Baqarah 2:19', 'ٱلْمَوْتِ'): ('death', r'dread of (\[death\])|dread of (\[the\]) death|dread (\[of\]) death', r'dread of ([death])'),
    ('An-Nisa 4:78', 'يَفْقَهُونَ'): ('understand', r'hardly (\[understand\]) any statement|hardly understand (\[to\]) any statement|they hardly (\[understand\])', r'they hardly ([understand]) any statement'),
    ('Al-Baqarah 2:44', 'تَعْقِلُونَ'): ('use reason', r'will you not (\[use reason\])|Then do (\[you\]) not use reason', r'Then will you not ([use reason])'),
    ('Al-Baqarah 2:219', 'تَتَفَكَّرُونَ'): ('give thought', r'that you might (\[give thought\])|\(\[They\]\) ask you about wine', r'that you might ([give thought])'),
    ('An-Nisa 4:165', 'مُّبَشِّرِينَ'): ('bringers of good tidings', r'as (\[bringers of good tidings\])|bringers (\[of\]) good tidings', r'as ([bringers of good tidings])'),
    ('Al-Baqarah 2:213', 'مُنذِرِينَ'): ('warners', r'and (\[warners\])|bringers of good tidings (\[and\]) warners', r'bringers of good tidings and ([warners])'),
    ('Al-Baqarah 2:51', 'ٰعَدْنَا'): ('made an appointment', r'when We (\[made an appointment\])|when (\[We\]) made an appointment', r'when We ([made an appointment])'),
    ('Surah 85:22', 'مَّحْفُوظٍۭ'): ('Preserved', r'in a (\[Preserved\]) Slate|\[Inscribed\] (\[in\]) a Preserved Slate', r'[Inscribed] in a ([Preserved]) Slate'),
    ('Al-Baqarah 2:201', 'مِنْهُم'): {
        'target_ar': 'ٱلْـَٔاخِرَةِ',
        'v_ar': 'وَمِنْهُم مَّن يَقُولُ رَبَّنَآ ءَاتِنَا فِى ٱلدُّنْيَا حَسَنَةً وَفِى ([ٱلْـَٔاخِرَةِ]) حَسَنَةً وَقِنَا عَذَابَ ٱلنَّارِ',
        'target_en': 'the Hereafter',
        'sub': (r'and in (\[the Hereafter\])|good (\[and\]) in the Hereafter', r'and in ([the Hereafter]) [that which is] good')
    },
    ('Al-Baqarah 2:95', 'قَدَّمَتْ'): ('have put forth', r'hands (\[have put forth\])|because (\[of\]) what their hands have put forth', r'because of what their hands ([have put forth])'),
    ('Al-Baqarah 2:173', 'حَرَّمَ'): ('forbidden', r'has only (\[forbidden\]) to you|\(\[He\]\) has only forbidden to you', r'He has only ([forbidden]) to you'),
    ('Al-Baqarah 2:27', 'مَآ'): {
        'target_ar': 'أَمَرَ',
        'v_ar': 'ٱلَّذِينَ يَنقُضُونَ عَهْدَ ٱللَّهِ مِنۢ بَعْدِ مِيثَـٰقِهِۦ وَيَقْطَعُونَ مَآ ([أَمَرَ]) ٱللَّهُ بِهِۦٓ أَن يُوصَلَ وَيُفْسِدُونَ فِى ٱلْأَرْضِ ۚ أُو۟لَـٰٓئِكَ هُمُ ٱلْخَـٰسِرُونَ',
        'target_en': 'ordered',
        'sub': (r'Allāh has (\[ordered\]) to be joined|\(\[Who\]\) break the covenant', r'Allāh has ([ordered]) to be joined')
    },
    ('Surah 78:2', 'ٱلنَّبَإِ'): ('great news', r'About the (\[great news\])|About (\[the\]) great news', r'About the ([great news])'),
    ('Al-Baqarah 2:37', 'تَلَقَّىٰٓ'): {
        'target_ar': 'كَلِمَـٰتٍ',
        'v_ar': 'فَتَلَقَّىٰٓ ءَادَمُ مِن رَّبِّهِۦ ([كَلِمَـٰتٍ]) فَتَابَ عَلَيْهِ ۚ إِنَّهُۥ هُوَ ٱلتَّوَّابُ ٱلرَّحِيمُ',
        'target_en': 'words',
        'sub': (r'\[some\] (\[words\])|\(\[Then\]\) Adam received', r'from his Lord [some] ([words])')
    },
    ('Al-Baqarah 2:49', 'مِّنْ'): {
        'target_ar': 'نَجَّيْنَـٰكُم',
        'v_ar': 'وَإِذْ ([نَجَّيْنَـٰكُم]) مِّنْ ءَالِ فِرْعَوْنَ يَسُومُونَكُمْ سُوٓءَ ٱلْعَذَابِ يُذَبِّحُونَ أَبْنَآءَكُمْ وَيَسْتَحْيُونَ نِسَآءَكُمْ ۚ وَفِى ذَٰلِكُم بَلَآءٌ مِّن رَّبِّكُمْ عَظِيمٌ',
        'target_en': 'We saved you',
        'sub': (r'when (\[We saved you\])|saved you \[i\.e\., your forefathers\] (\[from\]) the people', r'when ([We saved you]) [i.e., your forefathers] from the people')
    },
    ('Al-Baqarah 2:97', 'بُشْرَىٰ'): ('good tidings', r'and (\[good tidings\]) for the believers|guidance (\[and\]) good tidings for the believers', r'guidance and ([good tidings]) for the believers'),
    ('Al-Baqarah 2:104', 'رَٰعِنَا'): ('Rāʿinā', r'"(\[Rāʿinā\])"|O (\[you\]) who have believed, say not.*?"Rāʿinā"', r'"([Rāʿinā])"'),
    ('Al-Baqarah 2:32', 'سُبْحَـٰنَكَ'): ('Exalted are You', r'"(\[Exalted are You\])|They said, "Exalted are (\[You\])', r'They said, "([Exalted are You]);'),
    ("Ali 'Imran 3:144", 'ٱلشَّـٰكِرِينَ'): ('grateful', r'reward the (\[grateful\])|reward (\[the\]) grateful', r'reward the ([grateful])'),
    ('An-Nisa 4:65', 'تَسْلِيمًا'): ('submission', r'in \[full\] (\[submission\])|submit (\[in\]) \[full\] submission', r'in [full] ([submission])'),
    ('Al-Baqarah 2:45', 'ٱلْخَـٰشِعِينَ'): ('humbly submissive', r'for the (\[humbly submissive\])|for (\[the\]) humbly submissive', r'for the ([humbly submissive])'),
    ('Maryam 19:73', 'نَدِيًّا'): ('assembly', r'better in (\[assembly\])|better (\[in\]) assembly', r'better in ([assembly])'),
    ('Al-Baqarah 2:53', 'ٱلْكِتَـٰبَ'): ('Scripture', r'Moses the (\[Scripture\])|Moses (\[the\]) Scripture', r'Moses the ([Scripture])'),
    ('Al-Hajj 22:40', 'صَلَوَٰتٌ'): ('synagogues', r'churches, (\[synagogues\])|churches, synagogues, (\[And\])', r'churches, ([synagogues])')
}

updated = 0
for r in range(2, ws.max_row + 1):
    cit = str(ws.cell(r, 10).value or '').strip()
    t_ar = str(ws.cell(r, 11).value or '').strip()
    
    rule = ISM_FIXES.get((cit, t_ar))
    if not rule:
        continue
        
    v_en = str(ws.cell(r, 14).value or '')
    
    if isinstance(rule, dict):
        if 'target_ar' in rule:
            ws.cell(r, 11).value = rule['target_ar']
        if 'v_ar' in rule:
            ws.cell(r, 12).value = rule['v_ar']
        target_en = rule['target_en']
        pat, rep = rule['sub']
    else:
        target_en, pat, rep = rule
        
    ws.cell(r, 13).value = target_en
    
    new_v_en = re.sub(pat, rep, v_en)
    if new_v_en == v_en:
        clean_v = re.sub(r'\(\[([^\]]+)\]\)', r'\1', v_en)
        # find target_en in clean_v
        m = re.search(r'\b(' + re.escape(target_en) + r')\b', clean_v, re.IGNORECASE)
        if m:
            s, e = m.span()
            new_v_en = clean_v[:s] + f'([{clean_v[s:e]}])' + clean_v[e:]
            
    ws.cell(r, 14).value = new_v_en
    updated += 1

print(f'Updated {updated} rows in Ism sheet.')
wb.save('db/qw_en.xlsx')
print('Saved db/qw_en.xlsx.')
