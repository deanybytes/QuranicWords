import openpyxl, re

wb = openpyxl.load_workbook('db/qw_en.xlsx')

def strip_tashkeel(text):
    t = str(text).replace('\u0670', 'ا').replace('\u06E5', 'و').replace('\u06E6', 'ي').replace('\u0640', '')
    t = re.sub(r'[\u064B-\u065F\u06D6-\u06ED\uFEFF]', '', t)
    t = re.sub(r'[إأآٱ]', 'ا', t)
    t = t.replace('ة', 'ه').replace('ى', 'ي')
    return t.strip()

def strip_prefix(t_ar, v_ar):
    m = re.match(r'^([وَفَ][\u064B-\u065F\u0670]?)(.+)$', t_ar)
    if not m:
        m = re.match(r'^([وف])(.+)$', t_ar)
    if m:
        pref = m.group(1)
        stem = m.group(2)
        target_bracketed = f'([{t_ar}])'
        if target_bracketed in v_ar:
            new_v_ar = v_ar.replace(target_bracketed, f'{pref}([{stem}])')
            return stem, new_v_ar
        else:
            return stem, v_ar
    return t_ar, v_ar

# 1. Clean Fil sheet
sheet_fil = wb['Fil']
fil_prefix_count = 0
for r in range(2, sheet_fil.max_row + 1):
    lemma = str(sheet_fil.cell(row=r, column=2).value or '')
    clean_lemma = strip_tashkeel(lemma)
    t_ar = str(sheet_fil.cell(row=r, column=11).value or '')
    clean_tar = strip_tashkeel(t_ar)
    v_ar = str(sheet_fil.cell(row=r, column=12).value or '')
    
    if (clean_tar.startswith('و') and not clean_lemma.startswith('و')) or \
       (clean_tar.startswith('ف') and not clean_lemma.startswith('ف')):
        stem, new_v = strip_prefix(t_ar, v_ar)
        sheet_fil.cell(row=r, column=11).value = stem
        sheet_fil.cell(row=r, column=12).value = new_v
        fil_prefix_count += 1

print(f'Cleaned {fil_prefix_count} Arabic prefix rows in Fil.')

# 2. Clean Ism sheet
sheet_ism = wb['Ism']
ism_prefix_count = 0
for r in range(2, sheet_ism.max_row + 1):
    lemma = str(sheet_ism.cell(row=r, column=2).value or '')
    clean_lemma = strip_tashkeel(lemma)
    t_ar = str(sheet_ism.cell(row=r, column=11).value or '')
    clean_tar = strip_tashkeel(t_ar)
    v_ar = str(sheet_ism.cell(row=r, column=12).value or '')
    
    if (clean_tar.startswith('و') and not clean_lemma.startswith('و')) or \
       (clean_tar.startswith('ف') and not clean_lemma.startswith('ف')):
        stem, new_v = strip_prefix(t_ar, v_ar)
        sheet_ism.cell(row=r, column=11).value = stem
        sheet_ism.cell(row=r, column=12).value = new_v
        ism_prefix_count += 1

print(f'Cleaned {ism_prefix_count} Arabic prefix rows in Ism.')

# Save updated qw_en.xlsx
wb.save('db/qw_en.xlsx')
print('Saved db/qw_en.xlsx successfully.')
