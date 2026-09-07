import openpyxl, json, re

with open('db/fil_canonical_1505.json') as f:
    fil_canon = json.load(f)
with open('db/ism_canonical_3091.json') as f:
    ism_canon = json.load(f)

wb = openpyxl.load_workbook('db/qw_ur.xlsx')

# 1. Fil
ws_fil = wb['Fil']
while ws_fil.max_row > 1506:
    ws_fil.delete_rows(ws_fil.max_row)

for i, row in enumerate(fil_canon):
    r = i + 2
    for c, key in [(1, 'Rank'), (2, 'Arabic Lemma'), (3, 'Transliteration'), (4, 'Root'), (5, 'POS'), (6, 'Occurrences'), (7, 'Coverage %'), (8, 'Sense Index')]:
        ws_fil.cell(r, c).value = row[key]
    for c, key in [(10, 'Verse Citation'), (11, 'Target Arabic Word'), (12, 'Full Arabic Verse')]:
        ws_fil.cell(r, c).value = row[key]

# 2. Ism
ws_ism = wb['Ism']
while ws_ism.max_row > 3092:
    ws_ism.delete_rows(ws_ism.max_row)

for i, row in enumerate(ism_canon):
    r = i + 2
    for c, key in [(1, 'Rank'), (2, 'Arabic Lemma'), (3, 'Transliteration'), (4, 'Root'), (5, 'POS'), (6, 'Occurrences'), (7, 'Coverage %'), (8, 'Sense Index')]:
        ws_ism.cell(r, c).value = row[key]
    for c, key in [(10, 'Verse Citation'), (11, 'Target Arabic Word'), (12, 'Full Arabic Verse')]:
        ws_ism.cell(r, c).value = row[key]

wb.save('db/qw_ur.xlsx')
print('qw_ur.xlsx aligned. Fil max_row:', ws_fil.max_row, 'Ism max_row:', ws_ism.max_row)
