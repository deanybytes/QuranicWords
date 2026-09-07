import openpyxl, json, re, glob, os

workbooks = [
    ('en', 'db/qw_en.xlsx'),
    ('bn', 'db/qw_bn.xlsx'),
    ('ur', 'db/qw_ur.xlsx'),
    ('hi', 'db/qw_hi.xlsx'),
    ('id', 'db/qw_id.xlsx'),
    ('ms', 'db/qw_ms.xlsx'),
    ('tr', 'db/qw_tr.xlsx'),
    ('fa', 'db/qw_fa.xlsx'),
    ('ha', 'db/qw_ha.xlsx'),
    ('sw', 'db/qw_sw.xlsx'),
]

canons = {
    'Harf': (json.load(open('db/harf_canonical_192.json')), 192),
    'Fil': (json.load(open('db/fil_canonical_1505.json')), 1505),
    'Ism': (json.load(open('db/ism_canonical_3091.json')), 3091),
}

all_clean = True

for lang, path in workbooks:
    if not os.path.exists(path):
        print(f"FAILED: File {path} not found!")
        all_clean = False
        continue
        
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"\n==================== Checking {lang.upper()} ({path}) ====================")
    
    for sname, (canon, exp_len) in canons.items():
        if sname not in wb.sheetnames:
            print(f"  [{sname}] ERROR: Sheet missing!")
            all_clean = False
            continue
            
        ws = wb[sname]
        actual_rows = ws.max_row - 1
        
        mismatches = 0
        unbracketed = 0
        missing_c13 = 0
        prefix_errs = 0
        
        if actual_rows != exp_len:
            print(f"  [{sname}] ERROR: Expected {exp_len} rows, got {actual_rows} rows!")
            all_clean = False
            
        for i, c in enumerate(canon):
            r = i + 2
            rank = ws.cell(r, 1).value
            lemma = str(ws.cell(r, 2).value or '')
            root = str(ws.cell(r, 4).value or '').replace(' ', '').replace('-', '')
            cit = ws.cell(r, 10).value
            tar_ar = str(ws.cell(r, 11).value or '')
            c13 = ws.cell(r, 13).value
            v_text = str(ws.cell(r, 14).value or '')
            
            if rank != c['Rank'] or cit != c['Verse Citation'] or tar_ar != c['Target Arabic Word']:
                mismatches += 1
                
            if not c13 or not str(c13).strip():
                missing_c13 += 1
                
            m = re.findall(r'\(\[([^\]]+)\]\)', v_text)
            if not m:
                unbracketed += 1
                
            # Prefix check on non-waw/fa lemmas:
            norm_lem = re.sub(r'[\u064B-\u065F\u0670\u06D6-\u06ED]', '', lemma).strip()
            norm_tar = re.sub(r'[\u064B-\u065F\u0670\u06D6-\u06ED]', '', tar_ar).strip()
            norm_root = re.sub(r'[\u064B-\u065F\u0670\u06D6-\u06ED]', '', root).strip()
            
            if norm_tar.startswith('و') and not norm_lem.startswith('و') and not norm_root.startswith('و'):
                prefix_errs += 1
            if norm_tar.startswith('ف') and not norm_lem.startswith('ف') and not norm_root.startswith('ف') and norm_tar not in ['فرعون']:
                prefix_errs += 1
                
        status = "PASSED" if (mismatches == 0 and unbracketed == 0 and missing_c13 == 0 and prefix_errs == 0 and actual_rows == exp_len) else "FAILED"
        if status != "PASSED":
            all_clean = False
        print(f"  [{sname}] {status} - Rows: {actual_rows}/{exp_len}, Canon Mismatch: {mismatches}, Unbracketed: {unbracketed}, Missing Col 13: {missing_c13}, Prefix Errs: {prefix_errs}")

if all_clean:
    print("\n=======================================================")
    print(">>> ALL 10 WORKBOOKS ARE 100% PERFECT AND VERIFIED! <<<")
    print("=======================================================")
else:
    print("\n>>> CRITICAL: SOME WORKBOOKS HAVE ERRORS! <<<")
