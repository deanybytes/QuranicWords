import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ur.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. HARF
ws_harf = wb['Harf']
harf_fixes = {
    7: ('تو', lambda v: clean_brackets(v).replace('خدا تو لوگوں', 'خدا ([تو]) لوگوں')),
    8: ('پر', lambda v: clean_brackets(v).replace('ہدایت پر ہیں', 'ہدایت ([پر]) ہیں')),
    32: ('نہ', lambda v: clean_brackets(v).replace('یا نہ کرو', 'یا ([نہ]) کرو')),
    72: ('وہ', lambda v: clean_brackets(v).replace('وہ تمہاری پوشاک ہیں', '([وہ]) تمہاری پوشاک ہیں')),
    94: ('عمل نہ کیا', lambda v: clean_brackets(v).replace('عمل نہ کیا', '([عمل نہ کیا])')),
    97: ('سبحان الله', lambda v: clean_brackets(v).replace('سبحان الله', '([سبحان الله])')),
    121: ('الزام نہ دے سکیں', lambda v: clean_brackets(v).replace('الزام نہ دے سکیں', '([الزام نہ دے سکیں])')),
    137: ('ان کے سامنے (کی طرف)', lambda v: clean_brackets(v).replace('ان کے سامنے (کی طرف)', '([ان کے سامنے (کی طرف)])')),
    139: ('ان کے پیچھے', lambda v: clean_brackets(v).replace('ان کے پیچھے', '([ان کے پیچھے])')),
    168: ('ان کو', lambda v: clean_brackets(v).replace('اور ان کو ہم', 'اور ([ان کو]) ہم')),
    178: ('کل', lambda v: clean_brackets(v).replace('گویا کل وہاں', 'گویا ([کل]) وہاں')),
}
for r, (target, fn) in harf_fixes.items():
    ws_harf.cell(r, 13).value = target
    ws_harf.cell(r, 14).value = fn(ws_harf.cell(r, 14).value)
print("Applied Urdu Harf fixes.")

# 2. FIL
ws_fil = wb['Fil']
fil_count = 0
for r in range(2, ws_fil.max_row + 1):
    c10 = ws_fil.cell(r, 10).value
    c11 = ws_fil.cell(r, 11).value
    if (c10, c11) == ('An-Nisa 4:3', 'تَعْدِلُوا۟'):
        v = str(ws_fil.cell(r, 14).value)
        v = v.replace('بارے([انصاف نہ کرسکوگے])', 'بارے ([انصاف نہ کرسکوگے])')
        ws_fil.cell(r, 14).value = v
        fil_count += 1
print(f"Applied {fil_count} Urdu Fil fixes.")

# 3. ISM
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:21', 'تَتَّقُونَ'): ('بچو', lambda v: clean_brackets(v).replace('بچو', '([بچو])')),
    ("Ali 'Imran 3:84", 'عِيسَىٰ'): ('عیسیٰ', lambda v: clean_brackets(v).replace('اور عیسیٰ', 'اور ([عیسیٰ])')),
    ('Al-Baqarah 2:41', 'كَافِرٍۭ'): ('منکرِ اول', lambda v: clean_brackets(v).replace('منکرِ اول', '([منکرِ اول])')),
    ('Al-Baqarah 2:26', 'كَثِيرًا'): ('بہتوں', lambda v: clean_brackets(v).replace('بہتوں', '([بہتوں])')),
    ("Ali 'Imran 3:96", 'مُبَارَكًا'): ('بابرکت', lambda v: clean_brackets(v).replace('بابرکت', '([بابرکت])')),
    ('Maryam 19:73', 'نَدِيًّا'): ('مجلسیں', lambda v: clean_brackets(v).replace('مجلسیں', '([مجلسیں])')),
    ('An-Nahl 16:83', 'يَعْرِفُونَ'): ('واقف ہیں', lambda v: clean_brackets(v).replace('واقف ہیں۔', '([واقف ہیں])۔')),
}
ism_count = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target, fn = ism_fixes[k]
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = fn(ws_ism.cell(r, 14).value)
        ism_count += 1
print(f"Applied {ism_count} Urdu Ism fixes.")

wb.save('db/qw_ur.xlsx')
print("db/qw_ur.xlsx saved successfully!")
