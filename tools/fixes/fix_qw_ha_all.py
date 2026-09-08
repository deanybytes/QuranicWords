import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ha.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. HARF
ws_harf = wb['Harf']
# Row 136: Yusuf 12:25 | لَدَا
ws_harf.cell(136, 13).value = 'a wurin ƙõfar'
v136 = clean_brackets(ws_harf.cell(136, 14).value)
ws_harf.cell(136, 14).value = v136.replace('a wurin ƙõfar', '([a wurin ƙõfar])')
print("Applied Hausa Harf fixes.")

# 2. FIL
ws_fil = wb['Fil']
fil_fixes = {
    ('Al-Baqarah 2:60', 'عَلِمَ'): ('sun san', lambda v: clean_brackets(v).replace('sun san', '([sun san])')),
    ('Al-Isra 17:9', 'يَهْدِى'): ('yanã shiryarwa', lambda v: clean_brackets(v).replace('yanã shiryarwa', '([yanã shiryarwa])')),
    ('Al-Baqarah 2:155', 'بَشِّرِ'): ('ka yi bishãra', lambda v: clean_brackets(v).replace('ka yi bishãra', '([ka yi bishãra])')),
    ('Ibrahim 14:13', 'أَوْحَىٰٓ'): ('Ya yi wahayi', lambda v: clean_brackets(v).replace('Ya yi wahayi', '([Ya yi wahayi])')),
    ('Yunus 10:109', 'يَحْكُمَ'): ('Ya yi hukunci', lambda v: clean_brackets(v).replace('Ya yi hukunci', '([Ya yi hukunci])')),
    ('Al-Isra 17:59', 'ءَاتَيْنَا'): ('Mun bai wa', lambda v: clean_brackets(v).replace('Mun bai wa', '([Mun bai wa])')),
    ('Al-Isra 17:54', 'يَشَأْ'): ('Ya so', lambda v: clean_brackets(v).replace('Idan Ya so', 'Idan ([Ya so])')),
    ('Al-Baqarah 2:196', 'ٱسْتَيْسَرَ'): ('ya sauƙaƙa', lambda v: clean_brackets(v).replace('ya sauƙaƙa', '([ya sauƙaƙa])')),
    ('Al-Baqarah 2:221', 'يَتَذَكَّرُونَ'): ('suna tunãwa', lambda v: clean_brackets(v).replace('suna tunãwa', '([suna tunãwa])')),
    ('Al-Baqarah 2:34', 'ٱسْتَكْبَرَ'): ('ya yi girman kai', lambda v: clean_brackets(v).replace('ya yi girman kai', '([ya yi girman kai])')),
    ("Ali 'Imran 3:144", 'ٱنقَلَبْتُمْ'): ('zã kujũya', lambda v: clean_brackets(v).replace('zã kujũya', '([zã kujũya])')),
}
fil_count = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target, fn = fil_fixes[k]
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = fn(ws_fil.cell(r, 14).value)
        fil_count += 1
print(f"Applied {fil_count} Hausa Fil fixes.")

# 3. ISM
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:131', 'رَبُّهُۥٓ'): ('Ubangijinsa', lambda v: clean_brackets(v).replace('Ubangijinsa', '([Ubangijinsa])')),
    ('Al-Baqarah 2:20', 'قَدِيرٌ'): ('Mai ĩkon yi', lambda v: clean_brackets(v).replace('Mai ĩkon yi', '([Mai ĩkon yi])')),
    ('Al-Baqarah 2:49', 'يَسُومُونَكُمْ'): ('su na taya muku', lambda v: clean_brackets(v).replace('su na taya muku', '([su na taya muku])')),
    ('An-Nisa 4:78', 'يَفْقَهُونَ'): ('fahimtar', lambda v: clean_brackets(v).replace('fahimtar', '([fahimtar])')),
    ('Al-Baqarah 2:171', 'نِدَآءً'): ('ƙãra', lambda v: clean_brackets(v).replace('ƙãra', '([ƙãra])')),
    ('Al-Baqarah 2:31', 'أَنۢبِـُٔونِى'): ('Ku gaya mini', lambda v: clean_brackets(v).replace('Ku gaya mini', '([Ku gaya mini])')),
    ('Al-Baqarah 2:118', 'ءَايَةٌ'): ('ãyã', lambda v: clean_brackets(v).replace('wata ãyã', 'wata ([ãyã])')),
    ('Al-Baqarah 2:48', 'عَدْلٌ'): ('fansa', lambda v: clean_brackets(v).replace('fansa', '([fansa])')),
    ('Al-Baqarah 2:34', 'ٱسْتَكْبَرَ'): ('ya yi girman kai', lambda v: clean_brackets(v).replace('ya yi girman kai', '([ya yi girman kai])')),
    ('Al-Baqarah 2:62', 'صَـٰلِحًا'): ('aikin ƙwarai', lambda v: clean_brackets(v).replace('aikin ƙwarai', '([aikin ƙwarai])')),
    ('Al-Baqarah 2:57', 'طَيِّبَـٰتِ'): ('mãsu dãɗin', lambda v: clean_brackets(v).replace('mãsu dãɗin', '([mãsu dãɗin])')),
    ('Al-Baqarah 2:48', 'يُنصَرُونَ'): ('taimakon su', lambda v: clean_brackets(v).replace('taimakon su', '([taimakon su])')),
    ('Al-Baqarah 2:57', 'ظَلَّلْنَا'): ('ya yi inuwa', lambda v: clean_brackets(v).replace('ya yi inuwa', '([ya yi inuwa])')),
    ("Ali 'Imran 3:137", 'سُنَنٌ'): ('misãlai', lambda v: clean_brackets(v).replace('misãlai', '([misãlai])')),
    ('Al-Baqarah 2:8', 'يَقُولُ'): ('cewa', lambda v: clean_brackets(v).replace('cewa:', '([cewa]):')),
    ('Al-Baqarah 2:26', 'بَعُوضَةً'): ('sauro', lambda v: clean_brackets(v).replace('sauro', '([sauro])')),
    ('An-Nahl 16:71', 'بَعْضَكُمْ'): ('sãshenku', lambda v: clean_brackets(v).replace('sãshenku', '([sãshenku])')),
    ('Al-Baqarah 2:129', 'ٱلْحِكْمَةَ'): ('hikimar', lambda v: clean_brackets(v).replace('hikimar', '([hikimar])')),
    ('Al-Baqarah 2:35', 'تَقْرَبَا'): ('ku kusanci', lambda v: clean_brackets(v).replace('ku kusanci', '([ku kusanci])')),
    ('Al-Baqarah 2:48', 'تَجْزِى'): ('Wadãtar', lambda v: clean_brackets(v).replace('Wadãtar', '([Wadãtar])')),
    ("Al-Ma'idah 5:78", 'عَصَوا۟'): ('sãɓãwar', lambda v: clean_brackets(v).replace('sãɓãwar', '([sãɓãwar])')),
    ('Al-Baqarah 2:32', 'سُبْحَـٰنَكَ'): ('Tsarki ya tabbata a gare Ka', lambda v: clean_brackets(v).replace('Tsarki ya tabbata a gare Ka', '([Tsarki ya tabbata a gare Ka])')),
}
ism_count = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target, fn = ism_fixes[k]
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = fn(ws_ism.cell(r, 14).value)
        ism_count += 1
print(f"Applied {ism_count} Hausa Ism fixes.")

wb.save('db/qw_ha.xlsx')
print("db/qw_ha.xlsx saved successfully!")
