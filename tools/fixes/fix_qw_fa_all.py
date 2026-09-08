import openpyxl, re

wb = openpyxl.load_workbook('db/qw_fa.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. HARF
ws_harf = wb['Harf']
harf_fixes = {
    41: ('نزد اوست', lambda v: clean_brackets(v).replace('نزد اوست', '([نزد اوست])')),
    44: ('اوست', lambda v: clean_brackets(v).replace('اوست که همه', '([اوست]) که همه')),
}
for r, (target, fn) in harf_fixes.items():
    ws_harf.cell(r, 13).value = target
    ws_harf.cell(r, 14).value = fn(ws_harf.cell(r, 14).value)
print("Applied Persian Harf fixes.")

# 2. FIL
ws_fil = wb['Fil']
fil_fixes = {
    ('Al-Baqarah 2:10', 'كَانُوا۟'): ('می‌گفتند', lambda v: clean_brackets(v).replace('که می‌گفتند', 'که ([می‌گفتند])')),
    ('An-Nisa 4:61', 'رَأَيْتَ'): ('می‌بینی', lambda v: clean_brackets(v).replace('منافقان را می‌بینی', 'منافقان را ([می‌بینی])')),
    ('Al-Furqan 25:42', 'يَرَوْنَ'): ('ببینند', lambda v: clean_brackets(v).replace('را ببینند،', 'را ([ببینند])،')),
    ('Al-Baqarah 2:10', 'يَكْذِبُونَ'): ('دروغ‌هایی', lambda v: clean_brackets(v).replace('دروغ‌هایی که', '([دروغ‌هایی]) که')),
    ('An-Nur 24:51', 'دُعُوٓا۟'): ('فراخوانده می‌شوند', lambda v: clean_brackets(v).replace('فراخوانده می‌شوند', '([فراخوانده می‌شوند])')),
    ("Ali 'Imran 3:37", 'وَجَدَ'): ('می‌یافت', lambda v: clean_brackets(v).replace('غذایی می‌یافت', 'غذایی ([می‌یافت])')),
    ('Al-Baqarah 2:14', 'لَقُوا۟'): ('ملاقات', lambda v: clean_brackets(v).replace('هنگام ملاقات', 'هنگام ([ملاقات])')),
    ("Ali 'Imran 3:193", 'يُنَادِى'): ('فرامی‌خوانَد', lambda v: clean_brackets(v).replace('فرامی‌خوانَد', '([فرامی‌خوانَد])')),
    ('Al-A\'raf 7:54', 'تَبَارَكَ'): ('پربرکت', lambda v: clean_brackets(v).replace('پربرکت', '([پربرکت])')),
    ("Ali 'Imran 3:144", 'ٱنقَلَبْتُمْ'): ('بازمی‌گردید', lambda v: clean_brackets(v).replace('بازمی‌گردید', '([بازمی‌گردید])')),
    ('Al-Baqarah 2:219', 'تَتَفَكَّرُونَ'): ('بیندیشید', lambda v: clean_brackets(v).replace('بیندیشید.', '([بیندیشید]).')),
    ('Al-Baqarah 2:52', 'تَشْكُرُونَ'): ('سپاسگزاری کنید', lambda v: clean_brackets(v).replace('سپاسگزاری کنید.', '([سپاسگزاری کنید]).')),
    ('Yunus 10:93', 'يَقْضِى'): ('داوری می‌کند', lambda v: clean_brackets(v).replace('داوری می‌کند.', '([داوری می‌کند]).')),
    ('Al-Baqarah 2:191', 'ٱقْتُلُوهُمْ'): ('بکشید', lambda v: clean_brackets(v).replace('یافتید، بکشید', 'یافتید، ([بکشید])')),
    ('Al-Baqarah 2:79', 'يَكْتُبُونَ'): ('می‌نویسند', lambda v: clean_brackets(v).replace('خود می‌نویسند،', 'خود ([می‌نویسند])،')),
    ('Al-Baqarah 2:17', 'تَرَكَهُمْ'): ('رهایشان کند', lambda v: clean_brackets(v).replace('رهایشان کند.', '([رهایشان کند]).')),
    ("Ali 'Imran 3:117", 'أَهْلَكَتْهُ'): ('نابودش می‌سازد', lambda v: clean_brackets(v).replace('نابودش می‌سازد.', '([نابودش می‌سازد]).')),
}
fil_count = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target, fn = fil_fixes[k]
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = fn(ws_fil.cell(r, 14).value)
        fil_count += 1
print(f"Applied {fil_count} Persian Fil fixes.")

# 3. ISM
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:118', 'ءَايَةٌ'): ('نشانه‌ای', lambda v: clean_brackets(v).replace('معجزه و نشانه‌ای', 'معجزه و ([نشانه‌ای])')),
    ('Al-Baqarah 2:106', 'ءَايَةٍ'): ('آیه‌ای', lambda v: clean_brackets(v).replace('هر آیه‌ای', 'هر ([آیه‌ای])')),
    ('Al-Baqarah 2:26', 'بَعُوضَةً'): ('پشه‌ای', lambda v: clean_brackets(v).replace('به پشه‌ای', 'به ([پشه‌ای])')),
    ("Ali 'Imran 3:24", 'دِينِهِم'): ('دینشان', lambda v: clean_brackets(v).replace('در دینشان', 'در ([دینشان])')),
    ('Al-Baqarah 2:52', 'تَشْكُرُونَ'): ('سپاسگزاری کنید', lambda v: clean_brackets(v).replace('سپاسگزاری کنید.', '([سپاسگزاری کنید]).')),
    ('Al-Fatihah 1:2', 'ٱلْحَمْدُ'): ('سپاس‌ها', lambda v: clean_brackets(v).replace('سپاس‌ها]', '([سپاس‌ها])]')),
    ('Al-Baqarah 2:211', 'نِعْمَةَ'): ('نعمت', lambda v: clean_brackets(v).replace('نعمتِ الله', '([نعمتِ]) الله')),
    ("Ali 'Imran 3:33", 'إِبْرَٰهِيمَ'): ('ابراهیم', lambda v: clean_brackets(v).replace('خاندان \u200cابراهیم', 'خاندان ([ابراهیم])').replace('خاندان ‌ابراهیم', 'خاندان ([ابراهیم])').replace('خاندان ابراهیم', 'خاندان ([ابراهیم])')),
    ("Ali 'Imran 3:33", 'نُوحًا'): ('نوح', lambda v: clean_brackets(v).replace('آدم و نوح', 'آدم و ([نوح])')),
    ('Al-Baqarah 2:7', 'سَمْعِهِمْ'): ('گوش‌هایشان', lambda v: clean_brackets(v).replace('گوش‌هایشان', '([گوش‌هایشان])')),
    ('Al-Baqarah 2:7', 'عَظِيمٌ'): ('بزرگی', lambda v: clean_brackets(v).replace('عذاب بزرگی', 'عذاب ([بزرگی])')),
    ('Al-Baqarah 2:129', 'ٱلْعَزِيزُ'): ('پیروزمند', lambda v: clean_brackets(v).replace('پیروزمندِ حکیمی', '([پیروزمندِ]) حکیمی')),
    ('Al-Baqarah 2:168', 'مُّبِينٌ'): ('آشکار', lambda v: clean_brackets(v).replace('دشمن آشکارِ', 'دشمن ([آشکارِ])')),
    ('Al-Baqarah 2:10', 'يَكْذِبُونَ'): ('دروغ‌هایی', lambda v: clean_brackets(v).replace('دروغ‌هایی که', '([دروغ‌هایی]) که')),
    ('An-Nahl 16:83', 'يَعْرِفُونَ'): ('می‌شناسند', lambda v: clean_brackets(v).replace('می‌شناسند؛', '([می‌شناسند])؛')),
    ('Al-Baqarah 2:174', 'بُطُونِهِمْ'): ('شکم‌هایشان', lambda v: clean_brackets(v).replace('شکم‌هایشان', '([شکم‌هایشان])')),
    ('Al-Baqarah 2:165', 'شَدِيدُ'): ('سخت‌کیفر', lambda v: clean_brackets(v).replace('سخت‌کیفر', '([سخت‌کیفر])')),
    ('Al-Baqarah 2:10', 'أَلِيمٌۢ'): ('دردناکی', lambda v: clean_brackets(v).replace('عذاب دردناکی', 'عذاب ([دردناکی])')),
    ("Ali 'Imran 3:96", 'مُبَارَكًا'): ('پربرکت', lambda v: clean_brackets(v).replace('پربرکت', '([پربرکت])')),
    ('Al-Baqarah 2:45', 'بِٱلصَّبْرِ'): ('شکیبایی', lambda v: clean_brackets(v).replace('شکیبایی', '([شکیبایی])')),
    ('Al-Baqarah 2:37', 'ٱلتَّوَّابُ'): ('توبه‌پذیر', lambda v: clean_brackets(v).replace('توبه‌پذیرِ', '([توبه‌پذیرِ])')),
    ('Al-Baqarah 2:117', 'بَدِيعُ'): ('پدیدآورندۀ', lambda v: clean_brackets(v).replace('پدیدآورندۀ', '([پدیدآورندۀ])')),
    ("Ali 'Imran 3:44", 'نُوحِيهِ'): ('وحی می‌کنیم', lambda v: clean_brackets(v).replace('وحی می‌کنیم.', '([وحی می‌کنیم]).')),
    ('Al-Baqarah 2:44', 'تَعْقِلُونَ'): ('نمی‌اندیشید', lambda v: clean_brackets(v).replace('نمی‌اندیشید؟', '([نمی‌اندیشید])؟')),
    ('Al-Baqarah 2:38', 'يَحْزَنُونَ'): ('اندوهگین می‌شوند', lambda v: clean_brackets(v).replace('اندوهگین می‌شوند»', '([اندوهگین می‌شوند])»')),
    ("Ali 'Imran 3:120", 'يَفْرَحُوا۟'): ('شادمان می‌شوند', lambda v: clean_brackets(v).replace('شادمان می‌شوند؛', '([شادمان می‌شوند])؛')),
    ('An-Nisa 4:165', 'مُّبَشِّرِينَ'): ('بشارت‌بخش', lambda v: clean_brackets(v).replace('بشارت‌بخش', '([بشارت‌بخش])')),
    ('Al-Baqarah 2:213', 'مُنذِرِينَ'): ('بیم‌دهنده', lambda v: clean_brackets(v).replace('بیم\u200cدهنده برانگیخت', '([بیم\u200cدهنده]) برانگیخت').replace('بیم‌دهنده برانگیخت', '([بیم‌دهنده]) برانگیخت')),
    ('Al-Baqarah 2:255', 'يَعْلَمُ'): ('می‌داند', lambda v: clean_brackets(v).replace('می‌داند و', '([می‌داند]) و')),
    ('An-Nisa 4:145', 'ٱلْأَسْفَلِ'): ('پایین‌ترین', lambda v: clean_brackets(v).replace('پایین‌ترین', '([پایین‌ترین])')),
    ('Al-Baqarah 2:95', 'قَدَّمَتْ'): ('پیش فرستاده‌اند', lambda v: clean_brackets(v).replace('پیش فرستاده‌اند', '([پیش فرستاده‌اند])')),
    ('Al-Anbya 21:2', 'مُّحْدَثٍ'): ('تازه‌ای', lambda v: clean_brackets(v).replace('تازه‌ای', '([تازه‌ای])')),
    ("Ali 'Imran 3:135", 'لِذُنُوبِهِمْ'): ('گناهانشان', lambda v: clean_brackets(v).replace('گناهانشان', '([گناهانشان])')),
    ('Al-Baqarah 2:19', 'ءَاذَانِهِم'): ('گوش‌هایشان', lambda v: clean_brackets(v).replace('گوش‌هایشان', '([گوش‌هایشان])')),
    ('Al-Baqarah 2:8', 'يَقُولُ'): ('می‌گویند', lambda v: clean_brackets(v).replace('می‌گویند:', '([می‌گویند]):')),
    ('Al-Hashr 59:24', 'ٱلْمُصَوِّرُ'): ('شکل‌دهندۀ', lambda v: clean_brackets(v).replace('شکل‌دهندۀ', '([شکل‌دهندۀ])')),
    ('Yunus 10:73', 'ٱلْمُنذَرِينَ'): ('هشداریافتگان', lambda v: clean_brackets(v).replace('سرانجامِ هشداریافتگان', 'سرانجامِ ([هشداریافتگان])')),
    ('At-Tawbah 9:114', 'مَّوْعِدَةٍ'): ('وعده‌ای', lambda v: clean_brackets(v).replace('وعده‌ای', '([وعده‌ای])').replace('وعده\u200cاى', '([وعده\u200cاى])')),
    ('Al-Baqarah 2:250', 'أَقْدَامَنَا'): ('گام‌هایمان', lambda v: clean_brackets(v).replace('گام‌هایمان', '([گام‌هایمان])')),
    ('Al-Hajj 22:40', 'صَلَوَٰتٌ'): ('عبادتگاه‌ها', lambda v: clean_brackets(v).replace('عبادتگاه‌ها[ی یهود]', '([عبادتگاه‌ها])[ی یهود]')),
    ("Ash-Shu'ara 26:193", 'ٱلرُّوحُ'): ('روح‌الامین', lambda v: clean_brackets(v).replace('روح‌ الامین', '([روح‌ الامین])').replace('روح\u200c الامین', '([روح\u200c الامین])')),
    ('Al-Baqarah 2:17', 'مَثَلُهُمْ'): ('داستان اینان', lambda v: clean_brackets(v).replace('داستان اینان', '([داستان اینان])')),
    ("Ali 'Imran 3:75", 'سَبِيلٌ'): ('گناهی', lambda v: clean_brackets(v).replace('گناهی بر ما نیست', '([گناهی]) بر ما نیست')),
}
ism_count = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target, fn = ism_fixes[k]
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = fn(ws_ism.cell(r, 14).value)
        ism_count += 1
print(f"Applied {ism_count} Persian Ism fixes.")

wb.save('db/qw_fa.xlsx')
print("db/qw_fa.xlsx saved successfully!")
