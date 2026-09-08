import openpyxl, re

wb = openpyxl.load_workbook('db/qw_tr.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. Harf
ws_harf = wb['Harf']
harf_fixes = {
    22: ('getirmekten', 'Şüphesiz Allah, bir sivrisineği ve ondan daha büyüğünü misal ([getirmekten]) çekinmez.'),
    76: ('size', 'Ey Peygamber! Eşlerine şöyle söyle: "Eğer dünya dirliğini ve süsünü istiyorsanız, gelin ([size]) bağışta bulunayım ve sizi güzellikle salıvereyim."'),
    111: ('neyi', 'Şüphesiz Allah, bir sivrisineği ve ondan daha büyüğünü misal getirmekten çekinmez. İnananlar bunun Rablerinden bir gerçek olduğunu bilirler. İnkar edenler ise "Allah bu misalle ([neyi]) murat etti?" derler, O, bu misalle birçoğunu saptırır, birçoğunu da yola getirir. Onunla saptırdığı yalnız fasıklardır.'),
    139: ('arkalarındakini', "Allah, O'ndan başka tanrı olmayan, kendisini uyuklama ve uyku tutmayan, Hayy, Kayyum olandır. Göklerde ve yerde olanlar O'nundur. İzni olmadan katında şefaat edecek kimdir? Onların önlerindekini ve ([arkalarındakini]) bilir."),
    154: ('o iki kişiye', "İçinizden fuhuş yapan ([o iki kişiye]) eziyet edin; tevbe edip ıslah olurlarsa onlardan vazgeçin. Doğrusu Allah tevbeleri kabul eden, Esirgeyen'dir."),
    180: ('sabahı', 'Azabımız onların sahasına indiğinde, uyarılanların ([sabahı]) ne kötü olur!'),
    182: ('kuşluk vaktidir', 'Musa: "Buluşma zamanımız sizin bayram gününüzde, insanların toplanacağı ([kuşluk vaktidir])" dedi.'),
    191: ('uygun bir yerde', 'Musa: "Biz de seninkine benzer bir büyü göstereceğiz; bizimle senin aranda bir vakit tayin et ki, ne senin ne bizim caymayacağımız düz ve ([uygun bir yerde]) buluşma vaktimiz olsun" dedi.')
}
for r, (c13, v) in harf_fixes.items():
    ws_harf.cell(r, 13).value = c13
    ws_harf.cell(r, 14).value = v
print("Applied Turkish Harf fixes.")

# 2. Fil
ws_fil = wb['Fil']
fil_fixes = {
    ('Al-Baqarah 2:255', 'شَآءَ'): ('dilediğinden', lambda v: v.replace('([dilediği])nden', '([dilediğinden])')),
    ('Al-Baqarah 2:186', 'دَعَانِ'): ('ettiğinde', lambda v: v.replace('([ettiği])nde', '([ettiğinde])')),
    ('Al-Baqarah 2:186', 'سَأَلَكَ'): ('sorarlarsa', lambda v: v.replace('([sorar])larsa', '([sorarlarsa])')),
    ('Al-Baqarah 2:30', 'نُسَبِّحُ'): ('ediyoruz', lambda v: v.replace('([ediyor])uz', '([ediyoruz])')),
}
fil_count = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target, fn = fil_fixes[k]
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = fn(str(ws_fil.cell(r, 14).value))
        fil_count += 1
print(f"Applied {fil_count} Turkish Fil fixes.")

# 3. Ism
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:106', 'ءَايَةٍ'): ('ayetin', lambda v: v.replace('([ayeti])n', '([ayetin])')),
    ('Ash-Shu\'ara 26:89', 'بِقَلْبٍ'): ('kalble', lambda v: clean_brackets(v).replace('kalble', '([kalble])')),
    ('Al-Baqarah 2:47', 'فَضَّلْتُكُمْ'): ('üstün kıldığımı', lambda v: clean_brackets(v).replace('üstün kıldığımı', '([üstün kıldığımı])')),
    ('Al-Baqarah 2:26', 'كَثِيرًا'): ('birçoğunu', lambda v: clean_brackets(v).replace('birçoğunu', '([birçoğunu])', 1)),
    ('Al-Baqarah 2:79', 'يَكْسِبُونَ'): ('kazandıklarına', lambda v: clean_brackets(v).replace('kazandıklarına!', '([kazandıklarına])!')),
    ('Al-Baqarah 2:38', 'خَوْفٌ'): ('korku', lambda v: clean_brackets(v).replace('korku', '([korku])')),
    ('Al-Baqarah 2:218', 'هَاجَرُوا۟'): ('edenler', lambda v: clean_brackets(v).replace('hicret edenler', 'hicret ([edenler])')),
    ('Al-Baqarah 2:188', 'بِٱلْإِثْمِ'): ('günaha', lambda v: v.replace('([günah])a', '([günaha])')),
    ('An-Nisa 4:65', 'تَسْلِيمًا'): ('tamamen', lambda v: v.replace('([tam])amen', '([tamamen])')),
    ('At-Tawbah 9:103', 'خُذْ'): ('al', lambda v: clean_brackets(v).replace('olarak al,', 'olarak ([al]),')),
}
ism_count = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target, fn = ism_fixes[k]
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = fn(str(ws_ism.cell(r, 14).value))
        ism_count += 1
print(f"Applied {ism_count} Turkish Ism fixes.")

wb.save('db/qw_tr.xlsx')
print("db/qw_tr.xlsx saved successfully!")
