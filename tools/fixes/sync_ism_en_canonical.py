import openpyxl, json

# 1. Update qw_en.xlsx
wb_en = openpyxl.load_workbook('db/qw_en.xlsx')
ws_en = wb_en['Ism']

en_fixes = {
    18: ('Al-Baqarah 2:22', 'وَٱلسَّمَآءَ', 'ٱلَّذِى جَعَلَ لَكُمُ ٱلْأَرْضَ فِرَٰشًا ([وَٱلسَّمَآءَ]) بِنَآءً وَأَنزَلَ مِنَ ٱلسَّمَآءِ مَآءً فَأَخْرَجَ بِهِۦ مِنَ ٱلثَّمَرَٰتِ رِزْقًا لَّكُمْ ۖ فَلَا تَجْعَلُوا۟ لِلَّهِ أَندَادًا وَأَنتُمْ تَعْلَمُونَ', 'the sky', 'the sky', '[He] who made for you the earth a bed [spread out] and ([the sky]) a ceiling and sent down from the sky, rain and brought forth thereby fruits as provision for you. So do not attribute to Allāh equals while you know [that there is nothing similar to Him].'),
    23: ('Al-Baqarah 2:106', 'شَىْءٍ', '۞ مَا نَنسَخْ مِنْ ءَايَةٍ أَوْ نُنسِهَا نَأْتِ بِخَيْرٍ مِّنْهَآ أَوْ مِثْلِهَآ ۗ أَلَمْ تَعْلَمْ أَنَّ ٱللَّهَ عَلَىٰ كُلِّ ([شَىْءٍ]) قَدِيرٌ', 'thing', 'things', 'We do not abrogate a verse or cause it to be forgotten except that We bring forth [one] better than it or similar to it. Do you not know that Allāh is over all ([things]) competent?'),
    33: ("Ali 'Imran 3:139", 'مُّؤْمِنِينَ', 'وَلَا تَهِنُوا۟ وَلَا تَحْزَنُوا۟ وَأَنتُمُ ٱلْأَعْلَوْنَ إِن كُنتُم ([مُّؤْمِنِينَ])', 'believer', 'believers', 'So do not weaken and do not grieve, and you will be superior if you are [true] ([believers]).'),
    38: ('Al-Baqarah 2:47', 'ٱلْعَـٰلَمِينَ', 'يَـٰبَنِىٓ إِسْرَٰٓءِيلَ ٱذْكُرُوا۟ نِعْمَتِىَ ٱلَّتِىٓ أَنْعَمْتُ عَلَيْكُمْ وَأَنِّى فَضَّلْتُكُمْ عَلَى ([ٱلْعَـٰلَمِينَ])', 'the worlds', 'the worlds', 'O Children of Israel, remember My favor that I have bestowed upon you and that I preferred you over ([the worlds]).'),
    39: ('Al-Baqarah 2:85', 'بِبَعْضِ', 'أَفَتُؤْمِنُونَ ([بِبَعْضِ]) ٱلْكِتَـٰبِ وَتَكْفُرُونَ بِبَعْضٍ', 'part of', 'part of', 'So do you believe in ([part of]) the Scripture and disbelieve in part?'),
    41: ('Al-Baqarah 2:35', 'ٱلظَّـٰلِمِينَ', 'وَلَا تَقْرَبَا هَـٰذِهِ ٱلشَّجَرَةَ فَتَكُونَا مِنَ ([ٱلظَّـٰلِمِينَ])', 'wrongdoer', 'the wrongdoers', 'And do not approach this tree, lest you be among ([the wrongdoers]).'),
    42: ('Al-Baqarah 2:193', 'ٱلظَّـٰلِمِينَ', 'فَإِنِ ٱنتَهَوْا۟ فَلَا عُدْوَٰنَ إِلَّا عَلَى ([ٱلظَّـٰلِمِينَ])', 'oppressor', 'the oppressors', 'And fight them until there is no fitnah and [until] the religion, [all and every aspect of it], is for Allāh. But if they cease, then there is to be no aggression except against ([the oppressors]).'),
    50: ('Al-Baqarah 2:23', 'عَبْدِنَا', 'وَإِن كُنتُمْ فِى رَيْبٍۢ مِّمَّا نَزَّلْنَا عَلَىٰ ([عَبْدِنَا]) فَأْتُوا۟ بِسُورَةٍۢ مِّن مِّثْلِهِۦ', 'servant', 'Our servant', 'And if you are in doubt about what We have sent down upon ([Our Servant] [i.e., Prophet Muḥammad]), then produce a surah the like thereof.')
}

for r, (cit, tar_ar, full_ar, mean_en, tar_en, full_en) in en_fixes.items():
    ws_en.cell(r, 9).value = mean_en
    ws_en.cell(r, 10).value = cit
    ws_en.cell(r, 11).value = tar_ar
    ws_en.cell(r, 12).value = full_ar
    ws_en.cell(r, 13).value = tar_en
    ws_en.cell(r, 14).value = full_en

# Fix the 28 rows pointing to ذَٰلِكَ
for r in range(2, ws_en.max_row + 1):
    if ws_en.cell(r, 11).value == 'ذَٰلِكَ':
        ws_en.cell(r, 9).value = 'Book'
        ws_en.cell(r, 11).value = 'ٱلْكِتَـٰبُ'
        ws_en.cell(r, 12).value = 'ذَٰلِكَ ([ٱلْكِتَـٰبُ]) لَا رَيْبَ ۛ فِيهِ ۛ هُدًۭى لِّلْمُتَّقِينَ'
        ws_en.cell(r, 13).value = 'Book'
        ws_en.cell(r, 14).value = 'This is the ([Book]) about which there is no doubt, a guidance for those conscious of Allāh -'

wb_en.save('db/qw_en.xlsx')
print('qw_en.xlsx updated.')

# 2. Sync to ism_canonical_3091.json
with open('db/ism_canonical_3091.json') as f:
    canon = json.load(f)

for idx, c in enumerate(canon):
    r = idx + 2
    c['English Meaning'] = ws_en.cell(r, 9).value
    c['Verse Citation'] = ws_en.cell(r, 10).value
    c['Target Arabic Word'] = ws_en.cell(r, 11).value
    c['Full Arabic Verse'] = ws_en.cell(r, 12).value
    c['Target Meaning English'] = ws_en.cell(r, 13).value
    c['Full English Verse'] = ws_en.cell(r, 14).value

with open('db/ism_canonical_3091.json', 'w', encoding='utf-8') as f:
    json.dump(canon, f, ensure_ascii=False, indent=2)

print('ism_canonical_3091.json updated.')
