import openpyxl, re

wb = openpyxl.load_workbook('db/qw_id.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. Harf
ws_harf = wb['Harf']
harf_fixes = {
    16: ('sampai', 'Dihalalkan bagimu pada malam hari puasa bercampur dengan istrimu. Mereka adalah pakaian bagimu, dan kamu adalah pakaian bagi mereka. Allah mengetahui bahwa kamu tidak dapat menahan dirimu sendiri, tetapi Dia menerima tobatmu dan memaafkan kamu. Maka sekarang campurilah mereka dan carilah apa yang telah ditetapkan Allah bagimu. Makan dan minumlah hingga jelas bagimu (perbedaan) antara benang putih dan benang hitam, yaitu fajar. Kemudian sempurnakanlah puasa itu ([sampai]) malam. Tetapi jangan kamu campuri mereka, ketika kamu beriktikaf dalam masjid. Itulah ketentuan Allah, maka janganlah kamu mendekatinya. Demikianlah Allah menerangkan ayat-ayat-Nya kepada manusia, agar mereka bertakwa.'),
    17: ('kepada', 'Allah Pelindung orang yang beriman. Dia mengeluarkan mereka dari kegelapan ([kepada]) cahaya. Dan orang-orang yang kafir, pelindung-pelindungnya adalah setan, yang mengeluarkan mereka dari cahaya kepada kegelapan. Mereka adalah penghuni neraka, mereka kekal di dalamnya.'),
    26: ('darinya', 'Dan takutlah kamu pada hari ketika tidak seorang pun dapat menggantikan (membela) orang lain sedikit pun, tebusan tidak diterima ([darinya]), syafaat tidak berguna baginya, dan mereka tidak akan ditolong.'),
    81: ('Itulah', '([Itulah]) umat yang telah lalu. Baginya apa yang telah mereka usahakan dan bagimu apa yang telah kamu usahakan. Dan kamu tidak akan diminta pertanggungjawaban tentang apa yang mereka kerjakan.'),
    140: ('di atasmu', 'Dan (ingatlah) ketika Kami mengambil janji kamu dan Kami angkat gunung (Sinai) ([di atasmu]) (seraya berfirman), "Pegang teguhlah apa yang telah Kami berikan kepadamu dan ingatlah apa yang ada di dalamnya, agar kamu bertakwa."'),
    141: ('di bawahnya', 'Dan sampaikanlah kabar gembira kepada orang-orang yang beriman dan berbuat kebajikan, bahwa untuk mereka (disediakan) surga-surga yang mengalir ([di bawahnya]) sungai-sungai. Setiap kali mereka diberi rezeki buah-buahan dari surga itu, mereka berkata, "Inilah rezeki yang diberikan kepada kami dahulu." Mereka telah diberi (buah-buahan) yang serupa. Dan di sana mereka (memperoleh) pasangan-pasangan yang disucikan. Mereka kekal di dalamnya.'),
    145: ('sekelilingnya', 'Perumpamaan mereka seperti orang-orang yang menyalakan api, setelah menerangi ([sekelilingnya]), Allah melenyapkan cahaya (yang menyinari) mereka dan membiarkan mereka dalam kegelapan, tidak dapat melihat.'),
    171: ('mulailah', 'Lalu setan membisikkan pikiran jahat kepada mereka agar menampakkan aurat mereka (yang selama ini) tertutup. Dan (setan) berkata, "Tuhanmu tidak melarang kamu mendekati pohon ini, melainkan supaya kamu berdua tidak menjadi malaikat atau tidak menjadi orang yang kekal (dalam surga)." Dan dia (setan) bersumpah kepada keduanya, "Sesungguhnya aku ini benar-benar termasuk pencari nasihat bagi kamu berdua." Maka setan membujuk keduanya dengan tipu daya. Ketika keduanya telah mencicipi (buah) pohon itu, tampaklah bagi keduanya auratnya, maka ([mulailah]) mereka menutupinya dengan daun-daun surga. Dan Tuhan menyeru mereka, "Bukankah Aku telah melarang kamu berdua dari pohon itu dan Aku katakan kepadamu bahwa sesungguhnya setan itu musuh yang nyata bagi kamu berdua?"'),
    179: ('besok', 'Dan jangan sekali-kali engkau mengatakan terhadap sesuatu, "Aku pasti melakukan itu ([besok]) pagi,"')
}

for r, (c13, v) in harf_fixes.items():
    ws_harf.cell(r, 13).value = c13
    ws_harf.cell(r, 14).value = v
print("Applied Indonesian Harf fixes.")

# 2. Ism
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:32', 'سُبْحَـٰنَكَ'): 'Mahasuci',
    ('Al-Baqarah 2:32', 'ٱلْحَكِيمُ'): 'Mahabijaksana',
    ('Al-Baqarah 2:63', 'قُوَّةٍ'): 'teguhlah',
    ('Al-Baqarah 2:27', 'ٱلَّذِينَ'): 'orang-orang',
    ('Al-Baqarah 2:259', 'يَوْمًا'): 'sehari',
    ("Ali 'Imran 3:36", 'مَرْيَمَ'): 'Maryam',
    ("Ali 'Imran 3:96", 'بِبَكَّةَ'): 'Bakkah',
    ('Al-Baqarah 2:30', 'أَعْلَمُ'): 'Aku mengetahui',
}

ism_applied = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target = ism_fixes[k]
        v_clean = clean_brackets(ws_ism.cell(r, 14).value)
        if k == ('Al-Baqarah 2:32', 'سُبْحَـٰنَكَ'):
            # replace Mahasuci
            v_new = v_clean.replace('Mahasuci', '([Mahasuci])', 1)
            ws_ism.cell(r, 13).value = 'Mahasuci'
            ws_ism.cell(r, 14).value = v_new
        elif k == ('Al-Baqarah 2:32', 'ٱلْحَكِيمُ'):
            v_new = v_clean.replace('Mahabijaksana', '([Mahabijaksana])', 1)
            ws_ism.cell(r, 13).value = 'Mahabijaksana'
            ws_ism.cell(r, 14).value = v_new
        elif k == ('Al-Baqarah 2:27', 'ٱلَّذِينَ'):
            # replace orang-orang
            v_new = re.sub(r'orang[–-]orang', '([orang-orang])', v_clean, count=1)
            ws_ism.cell(r, 13).value = 'orang-orang'
            ws_ism.cell(r, 14).value = v_new
        else:
            assert target in v_clean, f"Target '{target}' not in {v_clean}"
            ws_ism.cell(r, 13).value = target
            ws_ism.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        ism_applied += 1

print(f"Applied {ism_applied} Indonesian Ism fixes.")
wb.save('db/qw_id.xlsx')
print("db/qw_id.xlsx saved successfully!")
