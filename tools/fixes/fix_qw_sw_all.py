import openpyxl, re

wb = openpyxl.load_workbook('db/qw_sw.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. Fil
ws_fil = wb['Fil']
fil_fixes = {
    ('Al-Furqan 25:42', 'يَرَوْنَ'): 'watakapo iona',
}
fil_applied = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target = fil_fixes[k]
        v_clean = clean_brackets(ws_fil.cell(r, 14).value)
        assert target in v_clean, f"Fil: '{target}' not in {v_clean}"
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        fil_applied += 1
print(f"Applied {fil_applied} Swahili Fil fixes.")

# 2. Ism
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:202', 'ٱلْحِسَابِ'): 'kuhisabu',
    ('Al-Qasas 28:76', 'قَـٰرُونَ'): 'Qaruni',
    ('An-Nisa 4:31', 'كَرِيمًا'): 'patukufu',
    ('Al-Baqarah 2:48', 'يُنصَرُونَ'): 'hawatanusuriwa',
}
ism_applied = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target = ism_fixes[k]
        v_clean = clean_brackets(ws_ism.cell(r, 14).value)
        assert target in v_clean, f"Ism: '{target}' not in {v_clean}"
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        ism_applied += 1
print(f"Applied {ism_applied} Swahili Ism fixes.")

wb.save('db/qw_sw.xlsx')
print("db/qw_sw.xlsx saved successfully!")
