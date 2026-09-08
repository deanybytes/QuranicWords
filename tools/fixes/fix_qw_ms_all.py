import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ms.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. Harf
ws_harf = wb['Harf']
harf_fixes = {
    131: ('Ah', 'Dan Tuhanmu telah perintahkan, supaya engkau tidak menyembah melainkan kepadaNya semata-mata, dan hendaklah engkau berbuat baik kepada ibu bapa. Jika salah seorang dari keduanya, atau kedua-duanya sekali, sampai kepada umur tua dalam jagaan dan peliharaanmu, maka janganlah engkau berkata kepada mereka (sebarang perkataan kasar) sekalipun perkataan "([Ah])", dan janganlah engkau menengking mereka, tetapi katakanlah kepada mereka perkataan yang mulia (yang bersopan santun).'),
    141: ('di bawahnya', 'Dan berilah khabar gembira kepada orang-orang yang beriman dan beramal soleh, sesungguhnya mereka beroleh syurga yang mengalir ([di bawahnya]) beberapa sungai; tiap-tiap kali mereka diberikan satu pemberian dari sejenis buah-buahan syurga itu, mereka berkata: "Inilah yang telah diberikan kepada kami dahulu"; dan mereka diberikan rezeki itu yang sama rupanya (tetapi berlainan hakikatnya), dan disediakan untuk mereka dalam syurga itu pasangan-pasangan, isteri-isteri yang bersih suci, sedang mereka pula kekal di dalamnya selama-lamanya.'),
    171: ('mulailah', 'Maka syaitan mengelincirkan mereka berdua dari syurga itu dan menyebabkan mereka dikeluarkan dari nikmat yang mereka telah berada di dalamnya dan Kami berfirman: "Turunlah kamu! Sebahagian kamu menjadi musuh kepada sebahagian yang lain dan bagi kamu disediakan tempat kediaman di bumi, dan juga kemudahan hidup hingga ke suatu masa (mati)". Dan Kami berfirman pula: "Wahai Adam! Tinggallah engkau dan isterimu di dalam syurga serta makanlah dari makanannya sepuas-puasnya apa sahaja kamu berdua sukai, dan janganlah kamu hampiri pokok ini; (jika kamu menghampirinya) maka akan menjadilah kamu dari golongan orang-orang yang zalim". Dengan sebab itu dapatlah syaitan memujuk mereka berdua (memakan buah pohon itu) dengan tipu dayanya. Setelah mereka memakan (buah) pohon itu, terdedahlah kepada mereka aurat mereka dan ([mulailah]) keduanya menutup auratnya dengan daun-daun dari syurga.')
}
for r, (c13, v) in harf_fixes.items():
    ws_harf.cell(r, 13).value = c13
    ws_harf.cell(r, 14).value = v
print("Applied Malay Harf fixes.")

# 2. Fil
ws_fil = wb['Fil']
fil_fixes = {
    ('An-Nisa 4:61', 'رَأَيْتَ'): 'engkau melihat',
    ('Al-Baqarah 2:112', 'أَسْلَمَ'): 'menyerahkan dirinya',
    ('Al-Baqarah 2:196', 'ٱسْتَيْسَرَ'): 'mudah didapati',
    ('Al-Baqarah 2:30', 'نُسَبِّحُ'): 'bertasbih',
    ("Ali 'Imran 3:97", 'ٱسْتَطَاعَ'): 'mampu',
}
fil_applied = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target = fil_fixes[k]
        v_clean = clean_brackets(ws_fil.cell(r, 14).value)
        assert target in v_clean, f"Fil row {r} ({k}): '{target}' not in {v_clean}"
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        fil_applied += 1
print(f"Applied {fil_applied} Malay Fil fixes.")

# 3. Ism
ws_ism = wb['Ism']
ism_fixes = {
    ('Al-Baqarah 2:131', 'رَبُّهُۥٓ'): 'Tuhannya',
    ('An-Nur 24:35', 'نَارٌ'): 'api',
    ('Al-Baqarah 2:49', 'فِرْعَوْنَ'): 'Firaun',
    ('Taha 20:112', 'مَن'): 'sesiapa',
    ('Al-Baqarah 2:97', 'عَدُوًّا'): 'memusuhi',
}
ism_applied = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target = ism_fixes[k]
        v_clean = clean_brackets(ws_ism.cell(r, 14).value)
        if k == ('An-Nur 24:35', 'نَارٌ'):
            # In 24:35, replace 'api' at whole word boundary (not in 'tetapi')
            # 'walaupun ia tidak disentuh api'
            v_new = re.sub(r'\bapi\b', '([api])', v_clean, count=1)
            ws_ism.cell(r, 13).value = 'api'
            ws_ism.cell(r, 14).value = v_new
        else:
            assert target in v_clean, f"Ism row {r} ({k}): '{target}' not in {v_clean}"
            ws_ism.cell(r, 13).value = target
            ws_ism.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        ism_applied += 1
print(f"Applied {ism_applied} Malay Ism fixes.")

wb.save('db/qw_ms.xlsx')
print("db/qw_ms.xlsx saved successfully!")
