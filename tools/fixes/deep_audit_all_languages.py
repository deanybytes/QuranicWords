import openpyxl, re, json

workbooks = [
    ('en', 'db/qw_en.xlsx'),
    ('bn', 'db/qw_bn.xlsx'),
    ('ur', 'db/qw_ur.xlsx'),
    ('hi', 'db/qw_hi.xlsx'),
    ('id', 'db/qw_id.xlsx'),
    ('ms', 'db/qw_ms.xlsx'),
    ('tr', 'db/qw_tr.xlsx'),
    ('fa', 'db/qw_fa.xlsx'),
    ('ha', 'db/qw_ha.xlsx'),
    ('sw', 'db/qw_sw.xlsx'),
]

# Word boundary check for regex
# For non-Latin scripts (Arabic, Urdu, Persian, Hindi, Bengali), \b might not work on all boundaries,
# but we can check if the character before '([' and after '])' is a letter/word character.
LETTER_PATTERN = r'[^\s\(\)\[\].,;!?:/\"\'\-—।؛؟]'

for lang, path in workbooks:
    wb = openpyxl.load_workbook(path, data_only=True)
    stats = {'total': 0, 'intra_word': 0, 'unbracketed': 0, 'c13_mismatch': 0}
    sample_intra = []
    sample_mismatch = []
    
    for sname in ['Harf', 'Fil', 'Ism']:
        ws = wb[sname]
        for r in range(2, ws.max_row + 1):
            stats['total'] += 1
            cit = ws.cell(r, 10).value
            c13 = str(ws.cell(r, 13).value or '').strip()
            v_text = str(ws.cell(r, 14).value or '').strip()
            
            m = re.findall(r'\(\[([^\]]+)\]\)', v_text)
            if not m:
                stats['unbracketed'] += 1
            else:
                bracketed = m[0].strip()
                if c13 and bracketed != c13:
                    stats['c13_mismatch'] += 1
                    if len(sample_mismatch) < 3:
                        sample_mismatch.append((sname, r, cit, f"c13='{c13}' vs hl='{bracketed}'"))
                        
                # Check intra-word boundary
                # Look for char before ([ or char after ])
                m_intra = re.search(f'({LETTER_PATTERN})\(\[|\]\)({LETTER_PATTERN})', v_text)
                if m_intra:
                    stats['intra_word'] += 1
                    if len(sample_intra) < 3:
                        start_pos = max(0, m_intra.start() - 15)
                        end_pos = min(len(v_text), m_intra.end() + 15)
                        sample_intra.append((sname, r, cit, c13, v_text[start_pos:end_pos]))
                        
    print(f"=== {lang.upper()} (Total rows: {stats['total']}) ===")
    print(f"  Intra-word highlights: {stats['intra_word']}")
    print(f"  Unbracketed verses:   {stats['unbracketed']}")
    print(f"  Col 13 vs HL mismatch:{stats['c13_mismatch']}")
    if sample_intra:
        print("  Sample intra-word:", sample_intra)
    if sample_mismatch:
        print("  Sample mismatch:", sample_mismatch)

