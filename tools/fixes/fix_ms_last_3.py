import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ms.xlsx')
ws = wb['Ism']

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

fixes = {
    ("Ali 'Imran 3:36", 'مَرْيَمَ'): 'Maryam',
    ("Ali 'Imran 3:96", 'بِبَكَّةَ'): 'Makkah',
    ('Al-Baqarah 2:30', 'أَعْلَمُ'): 'Aku mengetahui',
}

applied = 0
for r in range(2, ws.max_row + 1):
    k = (ws.cell(r, 10).value, ws.cell(r, 11).value)
    if k in fixes:
        target = fixes[k]
        v_clean = clean_brackets(ws.cell(r, 14).value)
        assert target in v_clean, f"Target '{target}' not in {v_clean}"
        ws.cell(r, 13).value = target
        ws.cell(r, 14).value = v_clean.replace(target, f'([{target}])', 1)
        applied += 1

print(f"Applied {applied} fixes to Malay Ism.")
wb.save('db/qw_ms.xlsx')
print("db/qw_ms.xlsx saved successfully!")
