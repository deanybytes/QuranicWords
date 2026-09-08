import openpyxl, re

wb = openpyxl.load_workbook('db/qw_hi.xlsx')

def clean_brackets(text):
    return re.sub(r'\(\[([^\]]+)\]\)', r'\1', str(text or '')).strip()

# 1. HARF
ws_harf = wb['Harf']
harf_fixes = {
    3: ('में से', lambda v: clean_brackets(v).replace('उनमें से', '([उनमें से])')),
    5: ('में', lambda v: clean_brackets(v).replace('तुममें से', '([तुममें से])')),
    12: ('नहीं', lambda v: clean_brackets(v).replace('संदेह नहीं,', 'संदेह ([नहीं]),')),
    14: ('जो', lambda v: clean_brackets(v).replace('जो तुम्हारे पास', '([जो]) तुम्हारे पास')),
    15: ('न तो', lambda v: clean_brackets(v).replace('न तो छोड़ा है', '([न तो]) छोड़ा है').replace('न आपको छोड़ा', '([न]) आपको छोड़ा')),
    18: ('जो', lambda v: clean_brackets(v).replace('ऐसे हैं, जो कहते', 'ऐसे हैं, ([जो]) कहते')),
    21: ('कि', lambda v: clean_brackets(v).replace('दिए हैं, कि बात', 'दिए हैं, ([कि]) बात')),
    25: ('यह', lambda v: clean_brackets(v).replace('यह (क़ुरआन)', '([यह]) (क़ुरआन)')),
    27: ('से', lambda v: clean_brackets(v).replace('क़िबले से, जिस पर', 'क़िबले ([से]), जिस पर')),
    32: ('न', lambda v: clean_brackets(v).replace('या उन्हें न डराया', 'या उन्हें ([न]) डराया')),
    35: ('या', lambda v: clean_brackets(v).replace('पापी या कृतघ्न', 'पापी ([या]) कृतघ्न')),
    36: ('जब', lambda v: clean_brackets(v).replace('थे, जब याक़ूब', 'थे, ([जब]) याक़ूब')),
    39: ('काश', lambda v: clean_brackets(v).replace('काश वे जानते!', '([काश]) वे जानते!').replace('काश कि वे', '([काश]) कि वे')),
    41: ('के पास', lambda v: clean_brackets(v).replace('जिसके पास किताब', 'जिस([के पास]) किताब')),
    61: ('या', lambda v: clean_brackets(v).replace('हो या उन्हें', 'हो ([या]) उन्हें')),
    64: ('अल्लाह की क़सम', lambda v: clean_brackets(v).replace('अल्लाह की क़सम!', '([अल्लाह की क़सम])!')),
    82: ('इसमें', lambda v: clean_brackets(v).replace('इसमें तुम्हारे पालनहार', '([इसमें]) तुम्हारे पालनहार')),
    85: ('कब आएगी', lambda v: clean_brackets(v).replace('कब आएगी?', '([कब आएगी])?')),
    91: ('अपने पास से', lambda v: clean_brackets(v).replace('अपने पास से', '([अपने पास से])')),
    99: ('इसी प्रकार', lambda v: clean_brackets(v).replace('इसी प्रकार अल्लाह', '([इसी प्रकार]) अल्लाह')),
    104: ('विनाश है', lambda v: clean_brackets(v).replace('विनाश है नाप-तौल', '([विनाश है]) नाप-तौल')),
    107: ('रही बात', lambda v: clean_brackets(v).replace('रही बात नाव की', '([रही बात]) नाव की')),
    111: ('क्या', lambda v: clean_brackets(v).replace('क्या इरादा किया है?', '([क्या]) इरादा किया है?')),
    113: ('तुम जहाँ भी', lambda v: clean_brackets(v).replace('तुम जहाँ भी हो,', '([तुम जहाँ भी]) हो,')),
    116: ('केवल', lambda v: clean_brackets(v).replace('ज़कात तो केवल', 'ज़कात तो ([केवल])')),
    120: ('ताकि', lambda v: clean_brackets(v).replace('शोक दिया ताकि जो', 'शोक दिया ([ताकि]) जो')),
    121: ('ताकि', lambda v: clean_brackets(v).replace('फेर लो। ताकि लोगों', 'फेर लो। ([ताकि]) लोगों')),
    124: ('तो क्या', lambda v: clean_brackets(v).replace('तो क्या ये धरती', '([तो क्या]) ये धरती')),
    125: ('और क्या', lambda v: clean_brackets(v).replace('और क्या वे नहीं', '([और क्या]) वे नहीं')),
    128: ('लाओ', lambda v: clean_brackets(v).replace('कहो : लाओ', 'कहो : ([लाओ])').replace('कहो :लाओ', 'कहो : ([लाओ])').replace('कहो :  लाओ', 'कहो : ([लाओ])')),
    131: ('उफ़', lambda v: clean_brackets(v).replace("'उफ'़", "'([उफ़])'").replace("'उफ़'", "'([उफ़])'")),
    136: ('दरवाज़े के पास', lambda v: clean_brackets(v).replace('दरवाज़े के पास पाया।', '([दरवाज़े के पास]) पाया।')),
    139: ('उनके पीछे', lambda v: clean_brackets(v).replace('उनके पीछे', '([उनके पीछे])') if 'उनके पीछे' in clean_brackets(v) else clean_brackets(v).replace('न नींद', '([उनके पीछे]) न नींद')),
    140: ('ऊपर', lambda v: clean_brackets(v).replace('तुम्हारे ऊपर उठाया', 'तुम्हारे ([ऊपर]) उठाया')),
    145: ('आस-पास', lambda v: clean_brackets(v).replace('आस-पास की चीज़ों', '([आस-पास]) की चीज़ों')),
    151: ('किस चीज़ से', lambda v: clean_brackets(v).replace('किस चीज़ से पैदा', '([किस चीज़ से]) पैदा')),
    161: ('तुम दोनों', lambda v: clean_brackets(v).replace('तुम दोनों को प्रभुता', '([तुम दोनों]) को प्रभुता')),
    167: ('तुम', lambda v: clean_brackets(v).replace('हम या तुम', 'हम या ([तुम])')),
    168: ('उन्हें', lambda v: clean_brackets(v).replace('और उन्हें भी देंगे।', 'और ([उन्हें]) भी देंगे।')),
    169: ('क्यों', lambda v: clean_brackets(v).replace('फ़रिश्तों को क्यों नहीं', 'फ़रिश्तों को ([क्यों]) नहीं')),
    175: ('सिवाय इसके कि', lambda v: clean_brackets(v).replace('सिवाय इसके कि', '([सिवाय इसके कि])')),
    176: ('अत्यंत करुणा करने वाला', lambda v: clean_brackets(v).replace('अत्यंत करुणा करने वाला, अत्यंत दयावान् है।', '([अत्यंत करुणा करने वाला]), अत्यंत दयावान् है।')),
    177: ('कभी', lambda v: clean_brackets(v).replace('कामना कभी नहीं', 'कामना ([कभी]) नहीं')),
    178: ('कल', lambda v: clean_brackets(v).replace('जैसे वह कल थी', 'जैसे वह ([कल]) थी')),
    179: ('कल', lambda v: clean_brackets(v).replace('मैं इसे कल करने', 'मैं इसे ([कल]) करने')),
    184: ('रातों-रात', lambda v: clean_brackets(v).replace('रातों-रात', '([रातों-रात])')),
    188: ('बराबर', lambda v: clean_brackets(v).replace('उनपर बराबर है', 'उनपर ([बराबर]) है')),
    191: ('बराबर', lambda v: clean_brackets(v).replace('जो बराबर हो।', 'जो ([बराबर]) हो।')),
}
for r, (target, fn) in harf_fixes.items():
    ws_harf.cell(r, 13).value = target
    ws_harf.cell(r, 14).value = fn(ws_harf.cell(r, 14).value)
print("Applied Hindi Harf fixes.")

# 2. FIL
ws_fil = wb['Fil']
fil_fixes = {
    ('Al-Baqarah 2:8', 'ءَامَنَّا'): ('ईमान लाए', lambda v: clean_brackets(v).replace('दिन पर ईमान लाए,', 'दिन पर ([ईमान लाए]),')),
    ('Al-Baqarah 2:26', 'أَرَادَ'): ('इरादा किया', lambda v: clean_brackets(v).replace('क्या इरादा किया है?', 'क्या ([इरादा किया]) है?')),
    ('Al-Baqarah 2:48', 'يُؤْخَذُ'): ('लिया जाएगा', lambda v: clean_brackets(v).replace('फ़िदया (दंड राशि) लिया जाएगा', 'फ़िदया (दंड राशि) ([लिया जाएगा])')),
    ("Ali 'Imran 3:117", 'ظَلَمُوٓا۟'): ('अत्याचार किया', lambda v: clean_brackets(v).replace('जिन्होंने अपने ऊपर अत्याचार किया', 'जिन्होंने अपने ऊपर ([अत्याचार किया])')),
    ('Al-Baqarah 2:191', 'ٱقْتُلُوهُمْ'): ('क़त्ल करो', lambda v: clean_brackets(v).replace('और उन्हें क़त्ल करो', 'और उन्हें ([क़त्ल करो])')),
    ('Ibrahim 14:13', 'أَوْحَىٰٓ'): ('वह़्य की', lambda v: clean_brackets(v).replace('उनकी ओर वह़्य की कि', 'उनकी ओर ([वह़्य की]) कि')),
    ("Ali 'Imran 3:117", 'أَهْلَكَتْهُ'): ('नष्ट कर दे', lambda v: clean_brackets(v).replace('उसको नष्ट कर दे।', 'उसको ([नष्ट कर दे])।')),
    ('Al-Baqarah 2:195', 'أَحْسِنُوٓا۟'): ('नेकी करो', lambda v: clean_brackets(v).replace('न डालो तथा नेकी करो,', 'न डालो तथा ([नेकी करो]),')),
    ('Al-Baqarah 2:34', 'ٱسْتَكْبَرَ'): ('अभिमान किया', lambda v: clean_brackets(v).replace('इनकार किया और अभिमान किया', 'इनकार किया और ([अभिमान किया])')),
    ("Ali 'Imran 3:144", 'ٱنقَلَبْتُمْ'): ('फिर जाओगे', lambda v: clean_brackets(v).replace('एड़ियों के बल फिर जाओगे?', 'एड़ियों के बल ([फिर जाओगे])?')),
    ("Ali 'Imran 3:101", 'يَعْتَصِم'): ('मज़बूत थाम ले', lambda v: clean_brackets(v).replace('अल्लाह को मज़बूत थाम ले,', 'अल्लाह को ([मज़बूत थाम ले]),')),
    ("Ali 'Imran 3:49", 'أَنفُخُ'): ('फूँक मारता हूँ', lambda v: clean_brackets(v).replace('उसमें फूँक मारता हूँ,', 'उसमें ([फूँक मारता हूँ]),')),
    ('Al-Baqarah 2:175', 'أَصْبَرَهُمْ'): ('सब्र करने वाले हैं', lambda v: clean_brackets(v).replace('कितना अधिक सब्र करने वाले हैं?', 'कितना अधिक ([सब्र करने वाले हैं])?')),
    ('Al-Baqarah 2:48', 'يُنصَرُونَ'): ('मदद की जाएगी', lambda v: clean_brackets(v).replace('उनकी मदद की जाएगी।', 'उनकी ([मदद की जाएगी])।')),
    ('Al-Baqarah 2:200', 'قَضَيْتُم'): ('पूरे कर लो', lambda v: clean_brackets(v).replace('कार्य पूरे कर लो,', 'कार्य ([पूरे कर लो]),')),
}
fil_count = 0
for r in range(2, ws_fil.max_row + 1):
    k = (ws_fil.cell(r, 10).value, ws_fil.cell(r, 11).value)
    if k in fil_fixes:
        target, fn = fil_fixes[k]
        ws_fil.cell(r, 13).value = target
        ws_fil.cell(r, 14).value = fn(ws_fil.cell(r, 14).value)
        fil_count += 1
print(f"Applied {fil_count} Hindi Fil fixes.")

# 3. ISM
ws_ism = wb['Ism']

# Fix verse copying issues
v_2_64 = 'फिर इसके पश्चात भी तुम ([फिर गए]), तो यदि तुमपर अल्लाह की दया और उसकी कृपा न होती तो तुम अवश्य घाटे में पड़ जाते।'
for r in [226, 328]:
    ws_ism.cell(r, 13).value = 'फिर गए'
    ws_ism.cell(r, 14).value = v_2_64

v_2_76 = 'क्या तुम उन्हें वह बात ([बताते हो]) जो अल्लाह ने तुमपर खोली है, ताकि वे तुम्हारे रब के यहाँ इसके द्वारा तुमपर हुज्जत क़ायम करें?'
ws_ism.cell(270, 13).value = 'बताते हो'
ws_ism.cell(270, 14).value = v_2_76

v_2_23 = 'और यदि तुम उस (क़ुरआन) के विषय में किसी संदेह में हो, जो हमने अपने बंदे पर उतारा है, तो उसके जैसी एक सूरत ले आओ और अल्लाह के अतिरिक्त अपने गवाहों (सहायकों) को भी बुला लो, यदि तुम ([सच्चे]) हो।'

ism_fixes = {
    ('Surah 7:59', 'قَوْمِهِۦ'): ('उसकी जाति', lambda v: clean_brackets(v).replace('नूह़ को उसकी जाति की ओर', 'नूह़ को ([उसकी जाति]) की ओर')),
    ('Al-Baqarah 2:20', 'كُلِّ'): ('हर चीज़', lambda v: clean_brackets(v).replace('अल्लाह हर चीज़ पर', 'अल्लाह ([हर चीज़]) पर')),
    ('Al-Baqarah 2:2', 'ٱلْكِتَـٰبُ'): ('पुस्तक', lambda v: clean_brackets(v).replace('वह पुस्तक है,', 'वह ([पुस्तक]) है,')),
    ('Al-Baqarah 2:8', 'بِمُؤْمِنِينَ'): ('मोमिन', lambda v: clean_brackets(v).replace('हरगिज़ मोमिन नहीं।', 'हरगिज़ ([मोमिन]) नहीं।')),
    ('Al-Baqarah 2:17', 'ظُلُمَـٰتٍ'): ('अँधेरों', lambda v: clean_brackets(v).replace('तरह के अँधेरों में', 'तरह के ([अँधेरों]) में')),
    ('Al-Baqarah 2:17', 'نَارًا'): ('आग', lambda v: clean_brackets(v).replace('जिसने एक आग भड़काई,', 'जिसने एक ([आग]) भड़काई,')),
    ('An-Nur 24:35', 'نَارٌ'): ('आग', lambda v: clean_brackets(v).replace('यद्यपि उसे आग ने न', 'यद्यपि उसे ([आग]) ने न')),
    ('Al-Baqarah 2:86', 'ٱلدُّنْيَا'): ('सांसारिक जीवन', lambda v: clean_brackets(v).replace('बदले सांसारिक जीवन ख़रीद', 'बदले ([सांसारिक जीवन]) ख़रीद')),
    ('Al-Baqarah 2:157', 'رَحْمَةٌ'): ('दया', lambda v: clean_brackets(v).replace('तथा बड़ी दया है,', 'तथा बड़ी ([दया]) है,')),
    ('Al-Baqarah 2:153', 'ٱلصَّـٰبِرِينَ'): ('धैर्य रखने वालों', lambda v: clean_brackets(v).replace('अल्लाह धैर्य रखने वालों के', 'अल्लाह ([धैर्य रखने वालों]) के')),
    ('Al-Baqarah 2:202', 'ٱلْحِسَابِ'): ('ह़िसाब', lambda v: clean_brackets(v).replace('जल्द ह़िसाब लेने वाला', 'जल्द ([ह़िसाब]) लेने वाला')),
    ('Al-Baqarah 2:52', 'تَشْكُرُونَ'): ('शुक्रगुज़ार', lambda v: clean_brackets(v).replace('ताकि तुम शुक्रगुज़ार बनो।', 'ताकि तुम ([शुक्रगुज़ार]) बनो।')),
    ('Al-Baqarah 2:102', 'مُلْكِ'): ('राज्य', lambda v: clean_brackets(v).replace('सुलैमान के राज्य में', 'सुलैमान के ([राज्य]) में')),
    ('An-Nisa 4:101', 'ٱلصَّلَوٰةِ'): ('नमाज़', lambda v: clean_brackets(v).replace('तो नमाज़ क़स्र', 'तो ([नमाज़]) क़स्र')),
    ('Al-Baqarah 2:17', 'بِنُورِهِمْ'): ('उनका प्रकाश', lambda v: clean_brackets(v).replace('प्रकाशित कर दिया, तो अल्लाह ने उनका प्रकाश छीन लिया', 'प्रकाशित कर दिया, तो अल्लाह ने ([उनका प्रकाश]) छीन लिया')),
    ('Al-Isra 17:93', 'سُبْحَانَ'): ('पवित्र है', lambda v: clean_brackets(v).replace('मेरा पालनहार पवित्र है!', 'मेरा पालनहार ([पवित्र है])!')),
    ('An-Nisa 4:168', 'طَرِيقًا'): ('राह', lambda v: clean_brackets(v).replace('उन्हें कोई राह दिखा दे।', 'उन्हें कोई ([राह]) दिखा दे।')),
    ('An-Nisa 4:13', 'ٱلْفَوْزُ'): ('बड़ी सफलता', lambda v: clean_brackets(v).replace('यही बड़ी सफलता है।', 'यही ([बड़ी सफलता]) है।')),
    ("Ali 'Imran 3:33", 'نُوحًا'): ('नूह़', lambda v: clean_brackets(v).replace('आदम और नूह़ को', 'आदम और ([नूह़]) को')),
    ("Ali 'Imran 3:36", 'مَرْيَمَ'): ('मरयम', lambda v: clean_brackets(v).replace('उसका नाम मरयम रखा', 'उसका नाम ([मरयम]) रखा')),
    ('Al-Baqarah 2:49', 'فِرْعَوْنَ'): ('फ़िरऔनियों', lambda v: clean_brackets(v).replace('तुम्हें फ़िरऔनियों से मुक्ति', 'तुम्हें ([फ़िरऔनियों]) से मुक्ति')),
    ('Al-Baqarah 2:34', 'ٱسْتَكْبَرَ'): ('अभिमान किया', lambda v: clean_brackets(v).replace('इनकार किया और अभिमान किया', 'इनकार किया और ([अभिमान किया])')),
    ('Al-Baqarah 2:234', 'خَبِيرٌ'): ('अवगत', lambda v: clean_brackets(v).replace('पूरी तरह अवगत है।', 'पूरी तरह ([अवगत]) है।')),
    ('An-Nisa 4:31', 'كَرِيمًا'): ('प्रतिष्ठित', lambda v: clean_brackets(v).replace('तुम्हें प्रतिष्ठित स्थान', 'तुम्हें ([प्रतिष्ठित]) स्थान')),
    ('Hud 11:73', 'مَّجِيدٌ'): ('गौरवशाली', lambda v: clean_brackets(v).replace('अति प्रशंसित, गौरवशाली है।', 'अति प्रशंसित, ([गौरवशाली]) है।')),
    ('Al-Baqarah 2:168', 'مُّبِينٌ'): ('खुला शत्रु', lambda v: clean_brackets(v).replace('तुम्हारा खुला शत्रु है।', 'तुम्हारा ([खुला शत्रु]) है।')),
    ('Al-Baqarah 2:105', 'ٱلْمُشْرِكِينَ'): ('मुश्रिकों', lambda v: clean_brackets(v).replace('हों या मुश्रिकों (अनेकेश्वरवादियों)', 'हों या ([मुश्रिकों]) (अनेकेश्वरवादियों)')),
    ('Al-Baqarah 2:57', 'طَيِّبَـٰتِ'): ('पाक चीज़ों', lambda v: clean_brackets(v).replace('अच्छी पाक चीज़ों में से', 'अच्छी ([पाक चीज़ों]) में से')),
    ('Al-Baqarah 2:267', 'ٱلْخَبِيثَ'): ('रद्दी चीज़', lambda v: clean_brackets(v).replace('रद्दी चीज़ का इरादा', '([रद्दी चीज़]) का इरादा')),
    ('Al-Baqarah 2:41', 'قَلِيلًا'): ('थोड़ा मूल्य', lambda v: clean_brackets(v).replace('बदले थोड़ा मूल्य न लो', 'बदले ([थोड़ा मूल्य]) न लो')),
    ('Al-Baqarah 2:26', 'كَثِيرًا'): ('बहुतों को', lambda v: clean_brackets(v).replace('साथ बहुतों को गुमराह', 'साथ ([बहुतों को]) गुमराह')),
    ("Al-An'am 6:120", 'ظَـٰهِरَ'): ('खुले', lambda v: clean_brackets(v).replace('लोगो!) खुले पाप छोड़', 'लोगो!) ([खुले]) पाप छोड़')),
    ('Al-Baqarah 2:10', 'أَلِيمٌۢ'): ('दर्दनाक यातना', lambda v: clean_brackets(v).replace('लिए दर्दनाक यातना है,', 'लिए ([दर्दनाक यातना]) है,')),
    ('Al-Baqarah 2:90', 'مُّهِينٌ'): ('अपमानजनक', lambda v: clean_brackets(v).replace('काफ़िरों के लिए अपमानजनक', 'काफ़िरों के लिए ([अपमानजनक])')),
    ("Ali 'Imran 3:7", 'مُّحْكَمَـٰتٌ'): ('मोह़कम', lambda v: clean_brackets(v).replace('आयतें मोह़कम (स्पष्ट', 'आयतें ([मोह़कम]) (स्पष्ट')),
    ('Al-Baqarah 2:48', 'يُنصَرُونَ'): ('मदद की जाएगी', lambda v: clean_brackets(v).replace('उनकी मदद की जाएगी।', 'उनकी ([मदद की जाएगी])।')),
    ("Ali 'Imran 3:144", 'عَقِبَيْهِ'): ('एड़ियों के बल', lambda v: clean_brackets(v).replace('अपनी एड़ियों के बल फिर', 'अपनी ([एड़ियों के बल]) फिर')),
    ("Ali 'Imran 3:44", 'نُوحِيهِ'): ('वह़्य कर रहे हैं', lambda v: clean_brackets(v).replace('ओर वह़्य कर रहे हैं और', 'ओर ([वह़्य कर रहे हैं]) और')),
    ('Al-Baqarah 2:6', 'كَفَرُوا۟'): ('कुफ़्र किया', lambda v: clean_brackets(v).replace('लोगों ने कुफ़्र किया,', 'लोगों ने ([कुफ़्र किया]),')),
    ('An-Nisa 4:78', 'يَفْقَهُونَ'): ('बात समझने', lambda v: clean_brackets(v).replace('कोई बात समझने के क़रीब', 'कोई ([बात समझने]) के क़रीब')),
    ('Surah 85:22', 'مَّحْفُوظٍۭ'): ('लौहे महफ़ूज़', lambda v: clean_brackets(v).replace('लौहे मह़फ़ूज़ (सुरक्षित', '([लौहे महफ़ूज़]) (सुरक्षित').replace('लौह़े मह़फ़ूज़', '([लौहे महफ़ूज़])')),
    ('At-Tawbah 9:121', 'صَغِيرَةً'): ('थोड़ा', lambda v: clean_brackets(v).replace('वे थोड़ा या अधिक,', 'वे ([थोड़ा]) या अधिक,')),
    ('Al-Baqarah 2:95', 'قَدَّمَتْ'): ('आगे भेजा', lambda v: clean_brackets(v).replace('हाथों ने आगे भेजा और', 'हाथों ने ([आगे भेजा]) और')),
    ('Al-Baqarah 2:173', 'حَرَّमَ'): ('हराम की है', lambda v: clean_brackets(v).replace('चीज़ हराम की है, जिसपर', 'चीज़ ([हराम की है]), जिसपर')),
    ('Al-Baqarah 2:23', 'صَـٰدِقِينَ'): ('सच्चे', lambda v: v_2_23),
    ("Al-Ma'idah 5:48", 'شِرْعَةً'): ('शरीय़त', lambda v: clean_brackets(v).replace('लिए एक शरीय़त तथा', 'लिए एक ([शरीय़त]) तथा')),
    ("Ali 'Imran 3:137", 'سُنَنٌ'): ('परंपराएं', lambda v: clean_brackets(v).replace('समान परंपराएं (परिस्थितियाँ)', 'समान ([परंपराएं]) (परिस्थितियाँ)')),
    ('Al-Baqarah 2:2', 'ذَٰلِكَ'): ('यह', lambda v: clean_brackets(v).replace('यह (क़ुरआन) वह पुस्तक', '([यह]) (क़ुरआन) वह पुस्तक')),
    ('Ar-Rahman 55:29', 'شَأْنٍ'): ('कार्य', lambda v: clean_brackets(v).replace('एक (नए) कार्य में', 'एक (नए) ([कार्य]) में')),
    ('Surah 78:2', 'ٱلنَّبَإِ'): ('सूचना', lambda v: clean_brackets(v).replace('बड़ी सूचना के विषय', 'बड़ी ([सूचना]) के विषय')),
    ('Al-Baqarah 2:32', 'سُبْحَـٰنَكَ'): ('तू पवित्र है', lambda v: clean_brackets(v).replace('कहा : तू पवित्र है।', 'कहा : ([तू पवित्र है])।')),
    ('Al-Baqarah 2:6', 'ٱلَّذِينَ'): ('जिन लोगों ने', lambda v: clean_brackets(v).replace('निःसंदेह जिन लोगों ने कुफ़्र', 'निःसंदेह ([जिन लोगों ने]) कुफ़्र')),
    ('Al-Baqarah 2:238', 'حَـٰفِظُوا۟'): ('ध्यान रखो', lambda v: clean_brackets(v).replace('का ध्यान रखो तथा', 'का ([ध्यान रखो]) तथा')),
    ("Ali 'Imran 3:22", 'نَّـٰصِرِينَ'): ('मदद करने वाले', lambda v: clean_brackets(v).replace('उनकी मदद करने वाले कोई नहीं।', 'उनकी ([मदद करने वाले]) कोई नहीं।')),
    ("Ali 'Imran 3:144", 'ٱلشَّـٰكِرِينَ'): ('आभारियों को', lambda v: clean_brackets(v).replace('शीघ्र ही आभारियों को बदला', 'शीघ्र ही ([आभारियों को]) बदला')),
    ('Al-Baqarah 2:45', 'ٱلْخَـٰشِعِينَ'): ('समर्पण करने वालों', lambda v: clean_brackets(v).replace('पूर्ण समर्पण करने वालों पर', 'पूर्ण ([समर्पण करने वालों]) पर')),
    ('Maryam 19:73', 'نَدِيًّا'): ('सभा', lambda v: clean_brackets(v).replace('और सभा की दृष्टि से', 'और ([सभा]) की दृष्टि से')),
    ('At-Tawbah 9:103', 'خُذْ'): ('दान लें', lambda v: clean_brackets(v).replace('मालों में से दान लें, जिसके', 'मालों में से ([दान लें]), जिसके')),
    ('Ash-Shura 42:52', 'رُوحًا'): ('रूह़', lambda v: clean_brackets(v).replace('एक रूह़ (क़ुरआन)', 'एक ([रूह़]) (क़ुरआन)')),
    ("Ash-Shu'ara 26:193", 'ٱلرُّوحُ'): ('रूह़ुल-अमीन', lambda v: clean_brackets(v).replace('इसे रूह़ुल-अमीन (अत्यंत', 'इसे ([रूह़ुल-अमीन]) (अत्यंत')),
}
ism_count = 0
for r in range(2, ws_ism.max_row + 1):
    k = (ws_ism.cell(r, 10).value, ws_ism.cell(r, 11).value)
    if k in ism_fixes:
        target, fn = ism_fixes[k]
        ws_ism.cell(r, 13).value = target
        ws_ism.cell(r, 14).value = fn(ws_ism.cell(r, 14).value)
        ism_count += 1
print(f"Applied {ism_count} Hindi Ism fixes.")

wb.save('db/qw_hi.xlsx')
print("db/qw_hi.xlsx saved successfully!")
