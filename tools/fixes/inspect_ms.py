import openpyxl, re

wb = openpyxl.load_workbook('db/qw_ms.xlsx', data_only=True)
PUNCT = set(' \t\n\r()[]{}.,;:!?/\\\"\'‘’-—।؛؟<>«»')

def has_chop(v_text):
    for m in re.finditer(r'\(\[([^\]]+)\]\)', v_text):
        start, end = m.start(), m.end()
        char_before = v_text[start - 1] if start > 0 else ' '
        char_after = v_text[end] if end < len(v_text) else ' '
        if char_before not in PUNCT or char_after not in PUNCT:
            return True, char_before, m.group(1), char_after, v_text[max(0, start-15):min(len(v_text), end+15)]
    return False, None, None, None, None

for sname in ['Harf', 'Fil', 'Ism']:
    ws = wb[sname]
    chops = []
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, 14).value or '')
        c13 = str(ws.cell(r, 13).value or '').strip()
        cit = ws.cell(r, 10).value
        tar_ar = ws.cell(r, 11).value
        rank = ws.cell(r, 1).value
        flag, b, tok, a, ctx = has_chop(v)
        if flag:
            chops.append((r, rank, cit, tar_ar, c13, f'{b}[{tok}]{a}', ctx))
    print(f'Malay {sname} chops: {len(chops)}')
    for c in chops:
        print('  ', c)
